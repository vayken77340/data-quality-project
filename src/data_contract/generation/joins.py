"""Joins sheet: per-epic registry of pairwise join relationships.

The spec workbook ships one optional "joins" sheet (configured in
defaults.yaml under `joins.sheet_name`). Each row describes a single
join between two tables: source/target tables and columns, join type,
cardinality, business rule.

Output: a single `epics/<epic>/contracts/joins.yaml` per epic. The file
is independent of the per-table contracts but validated against them
(source/target tables must be generated tables; referenced columns
must exist as fields on the corresponding contracts).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl.workbook.workbook import Workbook

from dq_core._util import dump_yaml, now_iso_z
from data_contract.generation.config import JoinsSpec
from dq_core.contract import Contract, _Provenance
from dq_core.errors import RejectionError
from data_contract.generation.header_matcher import find_column, normalize
from data_contract.generation.sheet_io import (
    HEADER_SEARCH_DEPTH,
    cell as _cell,
    locate_header_row,
    row_is_empty as _row_is_empty,
)


# ---------------------------------------------------------------------------
# Normalization tables
# ---------------------------------------------------------------------------


# Canonical join types. The spec's value is normalized to one of these.
JOIN_TYPE_ALIASES: dict[str, str] = {
    # English
    "inner join": "INNER",
    "inner": "INNER",
    "join": "INNER",
    "left join": "LEFT",
    "left outer join": "LEFT",
    "left": "LEFT",
    "right join": "RIGHT",
    "right outer join": "RIGHT",
    "right": "RIGHT",
    "full join": "FULL",
    "full outer join": "FULL",
    "full": "FULL",
    "outer": "FULL",
    "cross join": "CROSS",
    "cross": "CROSS",
    # French
    "jointure interne": "INNER",
    "jointure gauche": "LEFT",
    "jointure droite": "RIGHT",
    "jointure complete": "FULL",
    "jointure complète": "FULL",
    "jointure externe complete": "FULL",
    "jointure externe complète": "FULL",
    "jointure croisée": "CROSS",
    "jointure croisee": "CROSS",
}


_CARDINALITY_RE = re.compile(
    r"^\s*(?P<left>1|n|m|\*|many)\s*(?:[:\->]+|to)\s*(?P<right>1|n|m|\*|many)\s*$",
    re.IGNORECASE,
)

_CARD_SIDE = r"(?:1|n|m|\*|many)"


def parse_cardinality(
    raw: str, separator: str | tuple[str, ...] | None = None,
) -> str | None:
    """Normalize a cardinality string to `1:1`, `1:n`, `n:1`, or `n:m`.

    With `separator=None`, accepts the permissive default set (`:`, `->`, ` to `).
    With an explicit separator (e.g. `"->"` or `("->", ":")`), only the declared
    dividers are accepted between the two sides (whitespace around them is still
    tolerated). Multi-character separators are matched greedy-longest-first.

    Both `n` and `m` are treated as "many"; the canonical output uses `n:m` for
    the many-to-many case. Returns `None` on unparseable input.
    """
    if separator is None:
        m = _CARDINALITY_RE.match(raw)
    else:
        seps = (separator,) if isinstance(separator, str) else separator
        ordered = sorted(seps, key=len, reverse=True)
        sep_group = "|".join(re.escape(s) for s in ordered)
        pattern = rf"^\s*(?P<left>{_CARD_SIDE})\s*(?:{sep_group})\s*(?P<right>{_CARD_SIDE})\s*$"
        m = re.match(pattern, raw, re.IGNORECASE)
    if m is None:
        return None
    left = "1" if m.group("left") == "1" else "n"
    right = "1" if m.group("right") == "1" else "n"
    if left == "1" and right == "1":
        return "1:1"
    if left == "1" and right == "n":
        return "1:n"
    if left == "n" and right == "1":
        return "n:1"
    return "n:m"


# ---------------------------------------------------------------------------
# Dataclasses
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class JoinRow:
    sheet_row: int
    source_table: str
    target_table: str
    source_column: str
    target_column: str
    join_type: str        # canonical: INNER / LEFT / RIGHT / FULL / CROSS
    cardinality: str | None
    comment: str | None
    description: str | None


@dataclass
class JoinsData:
    rows: list[JoinRow] = field(default_factory=list)
    errors: list[RejectionError] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return bool(self.errors)


@dataclass
class JoinsContract(_Provenance):
    # Inherits version / epic / generated_at / spec_file / spec_sheet / target
    # from _Provenance. Note: `target` has a default there, so `joins` needs
    # one too (Python's "non-default after default" rule).
    joins: list[JoinRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        out = self._provenance_dict()
        out["joins"] = [_join_row_to_dict(j) for j in self.joins]
        return out


@dataclass
class JoinsRejection(_Provenance):
    errors: list[RejectionError] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        out = self._provenance_dict()
        out["errors"] = [e.to_dict() for e in self.errors]
        return out


def _join_row_to_dict(j: JoinRow) -> dict[str, Any]:
    out: dict[str, Any] = {
        "source_table": j.source_table,
        "source_column": j.source_column,
        "target_table": j.target_table,
        "target_column": j.target_column,
        "type": j.join_type,
    }
    if j.cardinality is not None:
        out["cardinality"] = j.cardinality
    if j.description is not None:
        out["description"] = j.description
    if j.comment is not None:
        out["comment"] = j.comment
    return out


# ---------------------------------------------------------------------------
# Sheet reading
# ---------------------------------------------------------------------------


def read_joins_sheet(wb: Workbook, joins_spec: JoinsSpec) -> JoinsData:
    """Read the joins sheet end-to-end. Accumulates errors; a non-empty
    `errors` list means the data shouldn't be used to emit a clean joins file."""
    out = JoinsData()
    cm = joins_spec.column_mapping

    # Tolerant sheet lookup with the joins-specific rejection kinds.
    target = normalize(joins_spec.sheet_name)
    if not target:
        out.errors.append(RejectionError(
            kind="joins_sheet_not_found",
            message=f"joins sheet name {joins_spec.sheet_name!r} is empty",
        ))
        return out
    matches = [s for s in wb.sheetnames if normalize(s) == target]
    if not matches:
        out.errors.append(RejectionError(
            kind="joins_sheet_not_found",
            message=(
                f"joins sheet {joins_spec.sheet_name!r} not found in workbook "
                f"(visible sheets: {wb.sheetnames})"
            ),
        ))
        return out
    if len(matches) > 1:
        out.errors.append(RejectionError(
            kind="joins_sheet_ambiguous",
            message=f"joins sheet name {joins_spec.sheet_name!r} matches multiple sheets: {matches}",
        ))
        return out

    sheet_name = matches[0]
    ws = wb[sheet_name]

    # Header row (depth-5 scan). Only columns with `required: true` (default)
    # contribute to header discovery.
    _required_for_header = [
        cm.source_table, cm.target_table, cm.source_column, cm.target_column, cm.join_type,
    ]
    required_norm = {normalize(c.spec_name) for c in _required_for_header if c.column_required}
    located = locate_header_row(ws, required_norm)

    if located is None:
        out.errors.append(RejectionError(
            kind="header_not_found",
            message=(
                f"joins sheet {sheet_name!r}: could not locate a header row "
                f"with required columns {sorted(required_norm)} in the first {HEADER_SEARCH_DEPTH} rows"
            ),
        ))
        return out

    header_row, headers = located

    # Locate each declared column.
    indices: dict[str, int | None] = {
        "source_table":  find_column(headers, cm.source_table.spec_name),
        "target_table":  find_column(headers, cm.target_table.spec_name),
        "source_column": find_column(headers, cm.source_column.spec_name),
        "target_column": find_column(headers, cm.target_column.spec_name),
        "join_type":     find_column(headers, cm.join_type.spec_name),
        "cardinality":   find_column(headers, cm.cardinality.spec_name) if cm.cardinality else None,
        "comment":       find_column(headers, cm.comment.spec_name) if cm.comment else None,
        "description":   find_column(headers, cm.description.spec_name) if cm.description else None,
    }

    # Required column headers must be present.
    required_columns = {
        "source_table":  cm.source_table,
        "target_table":  cm.target_table,
        "source_column": cm.source_column,
        "target_column": cm.target_column,
        "join_type":     cm.join_type,
    }
    missing: list[str] = []
    for key, col in required_columns.items():
        if indices[key] is None and col.column_required:
            missing.append(key)
    if missing:
        out.errors.append(RejectionError(
            kind="header_not_found",
            message=f"joins sheet {sheet_name!r}: missing required column headers: {', '.join(missing)}",
        ))
        return out

    cell_indices = [i for i in indices.values() if i is not None]
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx <= header_row:
            continue
        if _row_is_empty(row, cell_indices):
            continue
        join_row, row_errors = _process_join_row(
            row, row_idx, sheet_name, cm, indices, required_columns,
        )
        out.errors.extend(row_errors)
        if join_row is not None:
            out.rows.append(join_row)

    return out


def _process_join_row(
    row: tuple,
    row_idx: int,
    sheet_name: str,
    cm: JoinsColumnMapping,
    indices: dict[str, int | None],
    required_columns: dict[str, object],
) -> tuple["JoinRow | None", list[RejectionError]]:
    """Validate + parse one joins-sheet row. Returns `(join_row, errors)`.

    Returns `(None, [])` for rows skipped because a column with
    `default_value` declared was blank (the spec author opted into "blanks
    are tolerable"). Returns `(None, errors)` when any per-cell validation
    fails. Otherwise returns a fully-formed `JoinRow` and the empty list.
    """
    raw_cells = {key: _cell(row, idx) for key, idx in indices.items()}
    errors: list[RejectionError] = []
    trimmed: dict[str, str] = {}
    skip_row = False

    # Per-row checks for each conceptually-required column. Blank cells in
    # a column WITHOUT `default_value` reject; blank in one WITH `default_value`
    # silently skips the row.
    for key, col in required_columns.items():
        if indices[key] is None:
            skip_row = True  # header was absent (column_required=false)
            continue
        value = raw_cells[key]
        if value is None or str(value).strip() == "":
            if not col.has_default:
                errors.append(RejectionError(
                    kind="missing_mandatory",
                    sheet_row=row_idx,
                    column=col.spec_name,
                    field=key,
                    message=(
                        f"joins sheet {sheet_name!r} row {row_idx}: "
                        f"column {col.spec_name!r} ({key}) requires a value but the cell is empty"
                    ),
                ))
            else:
                skip_row = True
            continue
        trimmed[key] = str(value).strip()

    if errors or skip_row:
        return None, errors

    canonical_type = JOIN_TYPE_ALIASES.get(trimmed["join_type"].lower())
    if canonical_type is None:
        errors.append(RejectionError(
            kind="invalid_join_type",
            sheet_row=row_idx,
            column=cm.join_type.spec_name,
            value=trimmed["join_type"],
            message=(
                f"joins sheet {sheet_name!r} row {row_idx}: "
                f"join type {trimmed['join_type']!r} is not recognized "
                f"(expected INNER / LEFT / RIGHT / FULL / CROSS or a known alias)"
            ),
        ))
        return None, errors

    cardinality_value: str | None = None
    raw_card = raw_cells.get("cardinality")
    if raw_card is not None and str(raw_card).strip() != "":
        cardinality_separator = cm.cardinality.separators if cm.cardinality else None
        parsed = parse_cardinality(str(raw_card), separator=cardinality_separator)
        if parsed is None:
            errors.append(RejectionError(
                kind="invalid_cardinality",
                sheet_row=row_idx,
                column=cm.cardinality.spec_name if cm.cardinality else None,
                value=raw_card,
                message=(
                    f"joins sheet {sheet_name!r} row {row_idx}: "
                    f"cardinality {raw_card!r} is not parseable "
                    f"(expected `1:1`, `1:n`, `n:1`, `n:m` or `1 -> n` style)"
                ),
            ))
            return None, errors
        cardinality_value = parsed

    return JoinRow(
        sheet_row=row_idx,
        source_table=trimmed["source_table"],
        target_table=trimmed["target_table"],
        source_column=trimmed["source_column"],
        target_column=trimmed["target_column"],
        join_type=canonical_type,
        cardinality=cardinality_value,
        comment=_trim_or_none(raw_cells.get("comment")),
        description=_trim_or_none(raw_cells.get("description")),
    ), errors


def validate_joins(rows: list[JoinRow], contracts_by_table: dict[str, Contract]) -> list[RejectionError]:
    """Verify each parsed join row points at known tables + existing columns.

    Returns the list of errors. An empty list means the rows are usable.
    """
    errors: list[RejectionError] = []
    for r in rows:
        for side, table_name, column_name in (
            ("source", r.source_table, r.source_column),
            ("target", r.target_table, r.target_column),
        ):
            contract = contracts_by_table.get(table_name)
            if contract is None:
                errors.append(RejectionError(
                    kind="unknown_join_table",
                    sheet_row=r.sheet_row,
                    field=f"{side}_table",
                    value=table_name,
                    message=(
                        f"joins sheet row {r.sheet_row}: {side} table {table_name!r} "
                        f"is not among the generated tables for this epic"
                    ),
                ))
                continue
            field_names = {f.silver_name for f in contract.fields}
            if column_name not in field_names:
                errors.append(RejectionError(
                    kind="unknown_join_column",
                    sheet_row=r.sheet_row,
                    field=f"{side}_column",
                    value=column_name,
                    message=(
                        f"joins sheet row {r.sheet_row}: {side} column {column_name!r} "
                        f"does not exist as a field on table {table_name!r}"
                    ),
                ))
    return errors


# ---------------------------------------------------------------------------
# Output writing
# ---------------------------------------------------------------------------


def write_joins_outputs(
    result: JoinsContract | JoinsRejection,
    contracts_dir: Path,
    *,
    write_history: bool = True,
) -> list[Path]:
    """Materialize the joins output, cleaning up the opposite side.

    On success: writes `joins.yaml` (+ `history/<version>/joins.yaml` when
    `write_history=True`), deletes `rejected/joins.yaml` if present.
    On rejection: writes `rejected/joins.yaml`, deletes `joins.yaml` if present.
    History is never touched on rejection.
    """
    canonical = contracts_dir / "joins.yaml"
    rejected = contracts_dir / "rejected" / "joins.yaml"

    if isinstance(result, JoinsContract):
        dump_yaml(canonical, result.to_dict())
        paths = [canonical]
        if write_history:
            history = contracts_dir / "history" / result.version / "joins.yaml"
            dump_yaml(history, result.to_dict())
            paths.append(history)
        if rejected.exists():
            rejected.unlink()
        return paths

    dump_yaml(rejected, result.to_dict())
    if canonical.exists():
        canonical.unlink()
    return [rejected]


def build_joins_result(
    *,
    version: str,
    epic: str,
    spec_file_rel: str,
    spec_sheet: str,
    joins_data: JoinsData,
    contracts_by_table: dict[str, Contract],
    target: str = "",
    now: str | None = None,
) -> JoinsContract | JoinsRejection:
    """Combine sheet-read errors with validation errors and produce the final
    artifact (clean contract or rejection)."""
    errors: list[RejectionError] = list(joins_data.errors)
    if not errors:
        errors.extend(validate_joins(joins_data.rows, contracts_by_table))

    generated_at = now if now is not None else now_iso_z()
    common = dict(
        version=version,
        epic=epic,
        generated_at=generated_at,
        spec_file=spec_file_rel,
        spec_sheet=spec_sheet,
        target=target,
    )
    if errors:
        return JoinsRejection(**common, errors=errors)
    return JoinsContract(**common, joins=list(joins_data.rows))


# ---------------------------------------------------------------------------
# Local helpers
# ---------------------------------------------------------------------------


def _trim_or_none(v: object | None) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None
