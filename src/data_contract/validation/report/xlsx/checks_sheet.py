"""Checks sheet: every enabled check + its description + emitted violation kind."""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from data_contract.validation.models import ValidationReport
from data_contract.validation.report.dimensions import violation_kind_for_check
from data_contract.validation.report.xlsx._common import (
    append, autosize, style_header_row,
)


def build(ws, report: ValidationReport, S) -> None:
    """List every enabled check with its YAML-supplied description and the
    violation kind it emits. Disabled checks are intentionally omitted."""
    headers = list(S.get("checks_sheet", "headers"))
    append(ws, headers)
    style_header_row(ws, ncols=len(headers))
    rm = report.run_metadata
    if rm is None:
        return
    for name in rm.checks_enabled:
        try:
            vk = violation_kind_for_check(name)
        except KeyError:
            vk = ""
        append(ws, [name, rm.checks_descriptions.get(name, ""), vk])
    if ws.max_row > 1:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    autosize(ws, ncols=len(headers), max_width=100)
