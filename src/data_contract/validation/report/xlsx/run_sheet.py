"""Run sheet: validation-level metadata (epic, status, target, checks, versions)."""

from __future__ import annotations

from openpyxl.styles import Font

from data_contract.validation.models import ValidationReport
from data_contract.validation.report.xlsx._common import append, autosize
from data_contract.validation.report.xlsx._styles import (
    FILL_GREEN, FILL_ORANGE, FILL_RED,
)


def build(ws, report: ValidationReport, S) -> None:
    rm = report.run_metadata
    L = S.get("run_sheet", "labels")
    no_target = S.get("run_sheet", "no_target")
    pass_label = S.get("status", "pass")
    fail_label = S.get("status", "fail")
    rows: list[tuple[str, object]] = [
        (L["epic"], report.epic),
        (L["generated_at"], report.generated_at),
        (L["status"], rm.status if rm else (fail_label if report.has_errors else pass_label)),
        (L["status_reason"], rm.status_reason if rm else ""),
        (L["duration_ms"], rm.duration_ms if rm else 0),
        (L["target"], (rm.target["name"] if rm and rm.target else no_target)),
        (L["checks_enabled"], ", ".join(rm.checks_enabled) if rm else ""),
        (L["checks_disabled"], ", ".join(rm.checks_disabled) if rm else ""),
        (L["cli_args"], " ".join(rm.cli_args) if rm else ""),
    ]
    bold = Font(bold=True)
    for label, value in rows:
        append(ws, [label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    status_row = 3
    status_cell = ws.cell(row=status_row, column=2)
    status_cell.fill = (
        FILL_RED if status_cell.value == fail_label
        else FILL_GREEN if status_cell.value == pass_label
        else FILL_ORANGE
    )

    append(ws, [])
    append(ws, [S.get("run_sheet", "contract_versions_heading")])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, italic=True)
    if rm:
        for tbl, ver in rm.contracts.items():
            append(ws, [tbl, ver])
    autosize(ws, ncols=2, max_width=80)
