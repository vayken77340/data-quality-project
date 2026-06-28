"""Shared openpyxl style constants for every XLSX sheet builder.

Centralised so the colour palette and bold-header treatment live in exactly
one place. New sheets import from here rather than re-instantiating Font /
PatternFill objects per module.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill

from dq_core.report.colors import hex_


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=hex_("header"))
HEADER_ALIGN = Alignment(vertical="center", horizontal="left")

FILL_RED = PatternFill("solid", fgColor=hex_("red"))
FILL_ORANGE = PatternFill("solid", fgColor=hex_("orange"))
FILL_GREEN = PatternFill("solid", fgColor=hex_("green"))


CHECK_STATUS_FILL = {
    "OK": FILL_GREEN,
    "WARNING": FILL_ORANGE,
    "ERROR": FILL_RED,
}


def severity_fill(severity: str | None) -> PatternFill | None:
    """The tint applied to severity columns. Info / unknown -> no fill."""
    if severity == "error":
        return FILL_RED
    if severity == "warning":
        return FILL_ORANGE
    return None
