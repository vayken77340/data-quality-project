"""Profile sheet: per-table field profile (type, format, null/distinct counts)."""

from __future__ import annotations

from openpyxl.styles import Font

from dq_core.report_models import ValidationReport
from dq_core.report.xlsx._common import (
    append, autosize, style_header_row,
)
from dq_core.report.xlsx._styles import (
    HEADER_ALIGN, HEADER_FILL,
)


def build(ws, report: ValidationReport, S) -> None:
    headers = list(S.get("profile_sheet", "headers"))
    pk_indicator = S.get("profile_sheet", "pk_indicator")
    fk_indicator = S.get("profile_sheet", "fk_indicator")
    section_font = Font(bold=True, color="FFFFFF", size=12)
    first_section = True
    for tr in report.table_reports:
        if tr.profile is None:
            continue
        if not first_section:
            append(ws, [])
        first_section = False

        append(ws, [S.fmt("profile_sheet", "table_heading_template", table=tr.table)])
        heading_row = ws.max_row
        ws.merge_cells(start_row=heading_row, start_column=1,
                       end_row=heading_row, end_column=len(headers))
        heading_cell = ws.cell(row=heading_row, column=1)
        heading_cell.font = section_font
        heading_cell.fill = HEADER_FILL
        heading_cell.alignment = HEADER_ALIGN

        append(ws, headers)
        style_header_row(ws, ncols=len(headers))

        # Null #, Null %, Distinct come from the metrics registry -- one source
        # of truth. Disabled metrics surface as "--" in their column.
        nc = tr.metrics["null_count"].values if "null_count" in tr.metrics else {}
        np_ = tr.metrics["null_percentage"].values if "null_percentage" in tr.metrics else {}
        dc = tr.metrics["distinct_count"].values if "distinct_count" in tr.metrics else {}
        for f in tr.profile.fields:
            null_count = nc.get(f.name)
            null_pct = np_.get(f.name)
            distinct_count = dc.get(f.name)
            append(ws, [
                f.name,
                f.type,
                f.type_format,
                pk_indicator if f.is_pk else "",
                fk_indicator if f.is_fk else "",
                null_count if null_count is not None else "--",
                f"{null_pct:.1f}" if null_pct is not None else "--",
                distinct_count if distinct_count is not None else "--",
                f.total,
            ])
    autosize(ws, ncols=len(headers), max_width=40)
