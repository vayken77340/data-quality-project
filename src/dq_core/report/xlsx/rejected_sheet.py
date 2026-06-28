"""Per-table rejected-rows sheet.

Layout: `Severity | Source File | <PK fields | OR Source Row> |
          <only the fields that actually have a violation in this table> |
          Check | Expected`.

When a row violates multiple checks, the Check and Expected cells stack one
line per violation (newline-joined) so a reader can map check N to expected N
at the same vertical offset. Field-block cells get the offending value for
that field only -- if the same field is hit by multiple checks on this row,
the values stack the same way.
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from dq_core.contract import Contract
from dq_core.report_models import TableReport
from dq_core.report.xlsx._common import (
    append, autosize, check_label, severity_rank, stringify_cell, style_header_row,
)
from dq_core.report.xlsx._styles import severity_fill


def build(ws, tr: TableReport, contract: Contract | None, S) -> None:
    H = S.get("rejected_sheet", "headers")

    # Restrict the contract-field block to fields that actually have a
    # violation in this table. Preserve the contract's declared field order
    # so columns stay stable across runs.
    violating_fields_set = {
        v.get("field") for r in tr.rejected_rows for v in r.violations if v.get("field")
    }
    all_contract_fields = list(contract.fields) if contract else []
    contract_fields = [f.name for f in all_contract_fields if f.name in violating_fields_set]

    field_type_label: dict[str, str] = {}
    if contract is not None:
        type_user_labels = S.get("rejected_sheet", "type_user_labels")
        for f in all_contract_fields:
            field_type_label[f.name] = type_user_labels.get(
                f.type.value, f.type.value
            )

    pk_fields = list(tr.pk_fields)
    has_pk = bool(pk_fields)
    key_headers = list(pk_fields) if has_pk else [H["source_row"]]

    headers: list[str] = (
        [H["severity"], H["source_file"], *key_headers]
        + list(contract_fields)
        + [H["check"], H["expected"]]
    )
    append(ws, headers)
    style_header_row(ws, ncols=len(headers))

    severity_col = 1
    # 1-based: severity (1) + source_file (1) + key columns + 1 to step past them.
    field_block_start = 1 + 1 + len(key_headers) + 1
    field_to_col = {
        name: field_block_start + idx for idx, name in enumerate(contract_fields)
    }
    n_field_cols = len(contract_fields)
    check_col = field_block_start + n_field_cols
    expected_col = check_col + 1

    wrap_align = Alignment(wrap_text=True, vertical="top")

    for r in tr.rejected_rows:
        # Errors first, then warnings, info; tie-break on field name. Check and
        # Expected stack in this order so readers can map line N to line N.
        ordered = sorted(
            r.violations,
            key=lambda v: (severity_rank(v.get("severity")), v.get("field") or ""),
        )
        worst_severity = r.worst_severity or (
            ordered[0].get("severity") if ordered else ""
        )

        row: list[Any] = [(worst_severity or "").upper(), r.source_file]
        if has_pk:
            for pk_name in pk_fields:
                row.append(stringify_cell(r.pk_values.get(pk_name)))
        else:
            row.append(r.source_row)

        field_block_values: list[list[str]] = [[] for _ in contract_fields]
        check_lines: list[str] = []
        expected_lines: list[str] = []
        for v in ordered:
            check_lines.append(check_label(v, S, field_type_label))
            expected_lines.append(v.get("expected", ""))
            field = v.get("field")
            if field and field in field_to_col:
                col_idx = field_to_col[field] - field_block_start
                raw = v.get("offending_value")
                rendered = "(null)" if raw is None else stringify_cell(raw)
                field_block_values[col_idx].append(rendered)

        row.extend("\n".join(values) for values in field_block_values)
        row.append("\n".join(check_lines))
        row.append("\n".join(expected_lines))
        append(ws, row)

        fill = severity_fill(worst_severity)
        if fill is not None:
            ws.cell(row=ws.max_row, column=severity_col).fill = fill

        excel_row = ws.max_row
        ws.cell(row=excel_row, column=check_col).alignment = wrap_align
        ws.cell(row=excel_row, column=expected_col).alignment = wrap_align
        for idx, values in enumerate(field_block_values):
            if len(values) > 1:
                ws.cell(
                    row=excel_row, column=field_block_start + idx,
                ).alignment = wrap_align

    if tr.rejected_rows_truncated:
        append(ws, [
            S.fmt("rejected_sheet", "truncated_template",
                  count=tr.rejected_rows_truncated),
        ])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True)

    if ws.max_row > 1:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    autosize(ws, ncols=len(headers), max_width=80)
