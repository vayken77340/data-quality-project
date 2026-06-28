"""Run issues sheet: schema gaps + operational failures (one row per issue)."""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from dq_core.report_models import ValidationReport
from dq_core.report.hints import HINTS
from dq_core.report.xlsx._common import (
    append, autosize, severity_rank, style_header_row, table_level_violations,
)
from dq_core.report.xlsx._styles import severity_fill


def build(ws, report: ValidationReport, S) -> None:
    """Populate the run-issues sheet (`run_issues_sheet` in the strings YAML).

    Headers: Severity / Table / Check / Field / Expected / Hint.
    The severity cell is colour-tinted to mirror the rejected sheet.
    """
    headers = list(S.get("run_issues_sheet", "headers"))
    field_placeholder = S.get("run_issues_sheet", "field_placeholder")
    labels = S.get("rejected_sheet", "check_labels")
    append(ws, headers)
    style_header_row(ws, ncols=len(headers))

    rows = sorted(
        table_level_violations(report),
        key=lambda v: (
            severity_rank(v.severity),
            v.table or "",
            v.kind,
            v.field or "",
        ),
    )

    for v in rows:
        try:
            hint = HINTS.get(v.kind, "")
        except Exception:
            hint = ""
        # No {type} placeholders apply to table-level kinds, so a plain lookup
        # is enough.
        check = labels.get(v.kind, v.kind)
        append(ws, [
            (v.severity or "").upper(),
            v.table,
            check,
            v.field or field_placeholder,
            v.expected,
            hint,
        ])
        fill = severity_fill(v.severity)
        if fill is not None:
            ws.cell(row=ws.max_row, column=1).fill = fill

    if ws.max_row > 1:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    autosize(ws, ncols=len(headers), max_width=80)
