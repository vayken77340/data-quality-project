"""Summary sheet: overall scorecard + per-table totals with inline dimensions."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Font

from dq_core.report_models import ValidationReport
from dq_core.report.aggregation import filtered_dimension_score
from dq_core.report.dimensions import compute_overall_score
from dq_core.report.xlsx._common import (
    append, autosize, style_header_row,
)
from dq_core.report.xlsx._styles import (
    FILL_GREEN, FILL_ORANGE, FILL_RED,
)


_COMPLETENESS_KINDS = frozenset({"nullable_violation", "column_missing"})
_PK_UNIQUENESS_KINDS = frozenset({"pk_not_unique"})
_FK_CONSISTENCY_KINDS = frozenset({"fk_not_found"})


def _has_fk_check(report: ValidationReport) -> bool:
    rm = report.run_metadata
    return bool(rm and "fk_existence" in rm.checks_enabled)


def build(ws, report: ValidationReport, S) -> None:
    counts = report.summary_counts
    overall = compute_overall_score(
        [(tr.total_rows, tr.score) for tr in report.table_reports if tr.score]
    )
    pass_label = S.get("status", "pass")
    fail_label = S.get("status", "fail")
    status = fail_label if report.has_errors else pass_label
    bold = Font(bold=True)

    OL = S.get("summary_sheet", "overall_labels")
    append(ws, [OL["score"], f"{overall.score:.1f}"])
    append(ws, [OL["status"], status])
    append(ws, [OL["errors"], counts["error"]])
    append(ws, [OL["warnings"], counts["warning"]])
    append(ws, [OL["info"], counts["info"]])
    for r in range(1, ws.max_row + 1):
        ws.cell(row=r, column=1).font = bold
    status_cell = ws.cell(row=2, column=2)
    status_cell.fill = FILL_RED if status == fail_label else FILL_GREEN

    append(ws, [])
    base_cols = list(S.get("summary_sheet", "scorecard_headers"))
    DL = S.get("summary_sheet", "dimension_labels")
    show_fk = _has_fk_check(report)
    dim_columns: list[tuple[str, frozenset]] = [
        (DL["completeness"], _COMPLETENESS_KINDS),
        (DL["pk_uniqueness"], _PK_UNIQUENESS_KINDS),
    ]
    if show_fk:
        dim_columns.append((DL["fk_consistency"], _FK_CONSISTENCY_KINDS))
    headers = base_cols + [name for name, _ in dim_columns]
    append(ws, headers)
    style_header_row(ws, ncols=len(headers))

    no_pk = S.get("summary_sheet", "no_pk_placeholder")

    sum_files = sum_total_rows = sum_clean = sum_err = sum_warn = sum_info = 0
    dim_weighted: dict[str, float] = {name: 0.0 for name, _ in dim_columns}
    overall_weighted_score = 0.0
    total_weight = 0

    for tr in report.table_reports:
        n_err = sum(1 for v in tr.violations if v.severity == "error")
        n_warn = sum(1 for v in tr.violations if v.severity == "warning")
        n_info = sum(1 for v in tr.violations if v.severity == "info")
        pass_rate = tr.score.score if tr.score else 100.0
        clean = tr.score.clean_rows if tr.score else tr.total_rows
        row: list[Any] = [
            tr.table,
            ", ".join(tr.pk_fields) if tr.pk_fields else no_pk,
            len(tr.input_files),
            clean,
            tr.total_rows,
            n_err,
            n_warn,
            n_info,
            f"{pass_rate:.1f}",
        ]
        for _, kinds in dim_columns:
            row.append(f"{filtered_dimension_score(tr.violations, kinds=kinds, total_rows=tr.total_rows):.1f}")
        append(ws, row)
        fill = FILL_RED if n_err else (FILL_ORANGE if n_warn else FILL_GREEN)
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).fill = fill

        sum_files += len(tr.input_files)
        sum_total_rows += tr.total_rows
        sum_clean += clean
        sum_err += n_err
        sum_warn += n_warn
        sum_info += n_info
        if tr.total_rows > 0:
            overall_weighted_score += pass_rate * tr.total_rows
            for name, kinds in dim_columns:
                dim_weighted[name] += filtered_dimension_score(
                    tr.violations, kinds=kinds, total_rows=tr.total_rows,
                ) * tr.total_rows
            total_weight += tr.total_rows

    overall_pass = overall_weighted_score / total_weight if total_weight else 0.0
    totals_label = S.get("summary_sheet", "totals_label")
    totals_row: list[Any] = [
        totals_label, "", sum_files, sum_clean, sum_total_rows,
        sum_err, sum_warn, sum_info, f"{overall_pass:.1f}",
    ]
    for name, _ in dim_columns:
        avg = dim_weighted[name] / total_weight if total_weight else 0.0
        totals_row.append(f"{avg:.1f}")
    append(ws, totals_row)
    for c in range(1, len(headers) + 1):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

    autosize(ws, ncols=len(headers), max_width=40)
