"""Coverage for the joins sheet (epic-level relationship registry).

Three layers:

1. Sheet-lookup / header / row-parsing  -> exercises read_joins_sheet.
2. Normalization                         -> parse_cardinality + JOIN_TYPE_ALIASES.
3. Cross-validation                      -> validate_joins (against built contracts).

Plus CLI-level integration for end-to-end joins.yaml emission.
"""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from data_contract.cli import main
from data_contract.generation.config import JoinsSpec
from data_contract.contract import Contract, FieldContract
from data_contract.errors import RejectionError
from data_contract.generation.joins import (
    JOIN_TYPE_ALIASES,
    JoinRow,
    JoinsContract,
    JoinsRejection,
    build_joins_result,
    parse_cardinality,
    read_joins_sheet,
    validate_joins,
)
from data_contract.type_mapping import Type

from tests.conftest import add_keys_sheet, minimal_defaults_yaml, write_test_parsers_yaml


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


_DEFAULT_JOINS_SPEC_YAML = """
sheet_name: Joins
column_mapping:
  source_table:  { spec_name: Source Table }
  target_table:  { spec_name: Target Table }
  source_column: { spec_name: Source Col }
  target_column: { spec_name: Target Col }
  join_type:     { spec_name: Type }
  cardinality:   { spec_name: Card,         default_value: null }
  comment:       { spec_name: Comment,      default_value: null }
  description:   { spec_name: Description,  default_value: null }
"""


_JOINS_HEADERS = ("Source Table", "Target Table", "Source Col", "Target Col", "Type", "Card", "Comment", "Description")


def _spec(yaml_str: str = _DEFAULT_JOINS_SPEC_YAML) -> JoinsSpec:
    return JoinsSpec.from_dict(yaml.safe_load(yaml_str))


def _wb_with_joins(rows: list[tuple], *, sheet_name: str = "Joins", header_row: int = 1) -> Workbook:
    wb = Workbook()
    wb.active.title = "Other"
    ws = wb.create_sheet(sheet_name)
    for _ in range(header_row - 1):
        ws.append([None] * len(_JOINS_HEADERS))
    ws.append(list(_JOINS_HEADERS))
    for row in rows:
        padded = list(row) + [None] * (len(_JOINS_HEADERS) - len(row))
        ws.append(padded[: len(_JOINS_HEADERS)])
    return wb


def _contract(table: str, fields: list[str]) -> Contract:
    return Contract(
        version="1.0",
        epic="E",
        generated_at="2026-06-02T14:00:00Z",
        spec_file="x.xlsx",
        spec_sheet=table,
        table=table,
        fields=[FieldContract(name=n, type=Type.STRING, nullable=True, description=None) for n in fields],
    )


# ---------------------------------------------------------------------------
# 1. Sheet-lookup layer
# ---------------------------------------------------------------------------


def test_joins_sheet_not_found():
    wb = Workbook()
    wb.active.title = "Only"
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "joins_sheet_not_found" for e in result.errors)


def test_joins_sheet_ambiguous():
    wb = Workbook()
    wb.active.title = "Joins"
    other = wb.create_sheet("joins ")
    other.append(["x"])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "joins_sheet_ambiguous" for e in result.errors)


# ---------------------------------------------------------------------------
# 2. Header layer
# ---------------------------------------------------------------------------


def test_joins_header_missing_required_column():
    wb = Workbook()
    ws = wb.create_sheet("Joins")
    ws.append(["Source Table", "Target Table"])  # missing required columns
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "header_not_found" for e in result.errors)


def test_joins_header_on_third_row():
    wb = _wb_with_joins(
        [("PROJECT", "PROJWBS", "proj_id", "proj_id", "LEFT JOIN", "1:n", None, None)],
        header_row=3,
    )
    result = read_joins_sheet(wb, _spec())
    assert not result.errors
    assert result.rows[0].sheet_row == 4


def test_joins_optional_columns_absent_ok():
    spec = _spec("""
sheet_name: Joins
column_mapping:
  source_table:  { spec_name: Source Table }
  target_table:  { spec_name: Target Table }
  source_column: { spec_name: Source Col }
  target_column: { spec_name: Target Col }
  join_type:     { spec_name: Type }
""")
    wb = Workbook()
    ws = wb.create_sheet("Joins")
    ws.append(["Source Table", "Target Table", "Source Col", "Target Col", "Type"])
    ws.append(["PROJECT", "PROJWBS", "proj_id", "proj_id", "LEFT JOIN"])
    result = read_joins_sheet(wb, spec)
    assert not result.errors
    assert result.rows[0].cardinality is None
    assert result.rows[0].comment is None
    assert result.rows[0].description is None


# ---------------------------------------------------------------------------
# 3. Row parsing
# ---------------------------------------------------------------------------


def test_joins_row_missing_mandatory_source_table():
    wb = _wb_with_joins([(None, "PROJWBS", "x", "x", "LEFT")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "missing_mandatory" and e.field == "source_table" for e in result.errors)


def test_joins_row_missing_mandatory_target_table():
    wb = _wb_with_joins([("PROJECT", None, "x", "x", "LEFT")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "missing_mandatory" and e.field == "target_table" for e in result.errors)


def test_joins_row_missing_mandatory_source_column():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", None, "x", "LEFT")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "missing_mandatory" and e.field == "source_column" for e in result.errors)


def test_joins_row_missing_mandatory_target_column():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", None, "LEFT")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "missing_mandatory" and e.field == "target_column" for e in result.errors)


def test_joins_row_missing_mandatory_join_type():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", None)])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "missing_mandatory" and e.field == "join_type" for e in result.errors)


# ---------------------------------------------------------------------------
# 4. Normalization
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("raw, expected", [
    ("LEFT JOIN", "LEFT"),
    ("left join", "LEFT"),
    ("LEFT", "LEFT"),
    ("JOINTURE GAUCHE", "LEFT"),
    ("INNER JOIN", "INNER"),
    ("inner", "INNER"),
    ("JOIN", "INNER"),
    ("RIGHT JOIN", "RIGHT"),
    ("FULL OUTER JOIN", "FULL"),
    ("CROSS JOIN", "CROSS"),
])
def test_join_type_normalization(raw, expected):
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", raw)])
    result = read_joins_sheet(wb, _spec())
    assert not result.errors
    assert result.rows[0].join_type == expected


def test_invalid_join_type():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", "FOOBAR JOIN")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "invalid_join_type" for e in result.errors)


@pytest.mark.parametrize("raw, expected", [
    ("1:1", "1:1"),
    ("1:n", "1:n"),
    ("1:N", "1:n"),
    ("1 -> n", "1:n"),
    ("1 ->n", "1:n"),
    ("1->n", "1:n"),
    ("1 to n", "1:n"),
    ("n:1", "n:1"),
    ("N to 1", "n:1"),
    ("n:m", "n:m"),
    ("n to n", "n:m"),
    ("* -> *", "n:m"),
    ("many to many", "n:m"),
])
def test_cardinality_normalization(raw, expected):
    assert parse_cardinality(raw) == expected


def test_invalid_cardinality_returns_none():
    assert parse_cardinality("manyish") is None
    assert parse_cardinality("") is None


def test_invalid_cardinality_row_rejects():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", "LEFT", "manyish")])
    result = read_joins_sheet(wb, _spec())
    assert any(e.kind == "invalid_cardinality" for e in result.errors)


def test_cardinality_blank_ok():
    wb = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", "LEFT", None)])
    result = read_joins_sheet(wb, _spec())
    assert not result.errors
    assert result.rows[0].cardinality is None


# ---------------------------------------------------------------------------
# 5. Cross-validation
# ---------------------------------------------------------------------------


def test_unknown_source_table_rejects():
    rows = [JoinRow(2, "GHOST", "PROJWBS", "x", "x", "LEFT", None, None, None)]
    contracts = {"PROJWBS": _contract("PROJWBS", ["x"])}
    errors = validate_joins(rows, contracts)
    assert any(e.kind == "unknown_join_table" and e.field == "source_table" for e in errors)


def test_unknown_target_table_rejects():
    rows = [JoinRow(2, "PROJECT", "GHOST", "x", "x", "LEFT", None, None, None)]
    contracts = {"PROJECT": _contract("PROJECT", ["x"])}
    errors = validate_joins(rows, contracts)
    assert any(e.kind == "unknown_join_table" and e.field == "target_table" for e in errors)


def test_unknown_source_column_rejects():
    rows = [JoinRow(2, "PROJECT", "PROJWBS", "ghost", "x", "LEFT", None, None, None)]
    contracts = {"PROJECT": _contract("PROJECT", ["x"]), "PROJWBS": _contract("PROJWBS", ["x"])}
    errors = validate_joins(rows, contracts)
    assert any(e.kind == "unknown_join_column" and e.field == "source_column" for e in errors)


def test_unknown_target_column_rejects():
    rows = [JoinRow(2, "PROJECT", "PROJWBS", "x", "ghost", "LEFT", None, None, None)]
    contracts = {"PROJECT": _contract("PROJECT", ["x"]), "PROJWBS": _contract("PROJWBS", ["x"])}
    errors = validate_joins(rows, contracts)
    assert any(e.kind == "unknown_join_column" and e.field == "target_column" for e in errors)


def test_valid_join_passes_validation():
    rows = [JoinRow(2, "PROJECT", "PROJWBS", "proj_id", "proj_id", "LEFT", "1:n", None, "rule")]
    contracts = {
        "PROJECT": _contract("PROJECT", ["proj_id", "name"]),
        "PROJWBS": _contract("PROJWBS", ["proj_id", "wbs_name"]),
    }
    assert validate_joins(rows, contracts) == []


# ---------------------------------------------------------------------------
# 6. build_joins_result + output dataclasses
# ---------------------------------------------------------------------------


def test_build_joins_result_success():
    rows = [JoinRow(2, "PROJECT", "PROJWBS", "proj_id", "proj_id", "LEFT", "1:n", None, "rule")]
    contracts = {
        "PROJECT": _contract("PROJECT", ["proj_id"]),
        "PROJWBS": _contract("PROJWBS", ["proj_id"]),
    }
    from data_contract.generation.joins import JoinsData
    result = build_joins_result(
        version="1.0",
        epic="E",
        spec_file_rel="x.xlsx",
        spec_sheet="Joins",
        joins_data=JoinsData(rows=rows),
        contracts_by_table=contracts,
        now="2026-06-02T14:00:00Z",
    )
    assert isinstance(result, JoinsContract)
    payload = result.to_dict()
    assert payload["joins"][0]["type"] == "LEFT"
    assert payload["joins"][0]["cardinality"] == "1:n"
    assert payload["joins"][0]["description"] == "rule"
    assert "comment" not in payload["joins"][0]


def test_build_joins_result_rejection_from_validation():
    rows = [JoinRow(2, "PROJECT", "GHOST", "x", "x", "LEFT", None, None, None)]
    from data_contract.generation.joins import JoinsData
    result = build_joins_result(
        version="1.0",
        epic="E",
        spec_file_rel="x.xlsx",
        spec_sheet="Joins",
        joins_data=JoinsData(rows=rows),
        contracts_by_table={"PROJECT": _contract("PROJECT", ["x"])},
        now="2026-06-02T14:00:00Z",
    )
    assert isinstance(result, JoinsRejection)
    assert any(e.kind == "unknown_join_table" for e in result.errors)


def test_build_joins_result_propagates_sheet_errors():
    from data_contract.generation.joins import JoinsData
    sheet_err = RejectionError(kind="joins_sheet_not_found", message="...")
    result = build_joins_result(
        version="1.0",
        epic="E",
        spec_file_rel="x.xlsx",
        spec_sheet="Joins",
        joins_data=JoinsData(errors=[sheet_err]),
        contracts_by_table={},
        now="2026-06-02T14:00:00Z",
    )
    assert isinstance(result, JoinsRejection)
    assert result.errors[0].kind == "joins_sheet_not_found"


# ---------------------------------------------------------------------------
# 7. Output file writing
# ---------------------------------------------------------------------------


def test_write_outputs_success_creates_canonical_and_history(tmp_path):
    from data_contract.generation.joins import write_joins_outputs
    contract = JoinsContract(
        version="1.0", epic="E", generated_at="t", spec_file="x.xlsx", spec_sheet="Joins",
        joins=[JoinRow(2, "A", "B", "x", "x", "LEFT", None, None, None)],
    )
    paths = write_joins_outputs(contract, tmp_path)
    canonical = tmp_path / "joins.yaml"
    history = tmp_path / "history" / "1.0" / "joins.yaml"
    assert canonical in paths and history in paths
    assert canonical.exists() and history.exists()


def test_write_outputs_success_deletes_stale_rejection(tmp_path):
    from data_contract.generation.joins import write_joins_outputs
    rejected = tmp_path / "rejected" / "joins.yaml"
    rejected.parent.mkdir(parents=True)
    rejected.write_text("stale\n", encoding="utf-8")
    contract = JoinsContract(
        version="1.0", epic="E", generated_at="t", spec_file="x.xlsx", spec_sheet="Joins", joins=[],
    )
    write_joins_outputs(contract, tmp_path)
    assert not rejected.exists()


def test_write_outputs_rejection_deletes_canonical_keeps_history(tmp_path):
    from data_contract.generation.joins import write_joins_outputs
    canonical = tmp_path / "joins.yaml"
    canonical.write_text("stale\n", encoding="utf-8")
    history_dir = tmp_path / "history" / "0.9"
    history_dir.mkdir(parents=True)
    history_file = history_dir / "joins.yaml"
    history_file.write_text("untouched\n", encoding="utf-8")

    rejection = JoinsRejection(
        version="1.0", epic="E", generated_at="t", spec_file="x.xlsx", spec_sheet="Joins",
        errors=[RejectionError(kind="joins_sheet_not_found", message="...")],
    )
    paths = write_joins_outputs(rejection, tmp_path)
    rej_path = tmp_path / "rejected" / "joins.yaml"
    assert paths == [rej_path]
    assert rej_path.exists()
    assert not canonical.exists()
    assert history_file.read_text(encoding="utf-8") == "untouched\n"


# ---------------------------------------------------------------------------
# 8. CLI integration
# ---------------------------------------------------------------------------


def _bootstrap_joins_epic(
    tmp_path: Path,
    repo_root: Path,
    *,
    include_joins_block: bool = True,
    joins_rows: list[tuple] | None = None,
) -> Path:
    (tmp_path / "configs").mkdir(parents=True, exist_ok=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    write_test_parsers_yaml(tmp_path / "configs")
    edir = tmp_path / "epics" / "E"
    (edir / "configs" / "contracts").mkdir(parents=True)
    (edir / "specs").mkdir(parents=True)
    (edir / "contracts").mkdir(parents=True)

    defaults = minimal_defaults_yaml()
    if include_joins_block:
        defaults += """
joins:
  sheet_name: Joins
  column_mapping:
    source_table:  { spec_name: Source Table }
    target_table:  { spec_name: Target Table }
    source_column: { spec_name: Source Col }
    target_column: { spec_name: Target Col }
    join_type:     { spec_name: Type }
    cardinality:   { spec_name: Card,         default_value: null }
    description:   { spec_name: Description,  default_value: null }
"""
    (tmp_path / "configs" / "specs_parsing.yaml").write_text(defaults, encoding="utf-8")
    (edir / "configs" / "contracts" / "v1.0.yaml").write_text(
        "epic: E\nversion: '1.0'\nspec_file_name: spec.xlsx\ntarget: postgres\n"
        "tables:\n  - table_name: T1\n  - table_name: T2\n",
        encoding="utf-8",
    )

    wb = Workbook()
    wb.active.title = "T1"
    t1 = wb["T1"]
    t1.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    t1.append(["a", "Double", "d", "OUI"])
    t2 = wb.create_sheet("T2")
    t2.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    t2.append(["a", "Double", "d", "OUI"])
    add_keys_sheet(wb, [("T1", "a"), ("T2", "a")])

    if include_joins_block:
        joins_ws = wb.create_sheet("Joins")
        joins_ws.append(["Source Table", "Target Table", "Source Col", "Target Col", "Type", "Card", "Description"])
        for r in (joins_rows or [("T1", "T2", "a", "a", "LEFT JOIN", "1:n", "demo")]):
            padded = list(r) + [None] * (7 - len(r))
            joins_ws.append(padded[:7])
    wb.save(edir / "specs" / "spec.xlsx")
    return edir


def test_cli_emits_joins_yaml_alongside_table_contracts(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_joins_epic(tmp_path, repo_root)
    rc = main(["generate", "--epic", "E"])
    assert rc == 0
    assert (edir / "contracts" / "T1.yaml").exists()
    assert (edir / "contracts" / "T2.yaml").exists()
    joins_file = edir / "contracts" / "joins.yaml"
    assert joins_file.exists()
    payload = yaml.safe_load(joins_file.read_text(encoding="utf-8"))
    assert payload["joins"][0]["source_table"] == "T1"
    assert payload["joins"][0]["type"] == "LEFT"
    assert payload["joins"][0]["cardinality"] == "1:n"
    history = edir / "contracts" / "history" / "1.0" / "joins.yaml"
    assert history.exists()


def test_cli_no_joins_block_no_joins_file(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_joins_epic(tmp_path, repo_root, include_joins_block=False)
    rc = main(["generate", "--epic", "E"])
    assert rc == 0
    assert (edir / "contracts" / "T1.yaml").exists()
    assert (edir / "contracts" / "T2.yaml").exists()
    assert not (edir / "contracts" / "joins.yaml").exists()


def test_cli_joins_rejection_when_referenced_column_missing(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_joins_epic(
        tmp_path, repo_root,
        joins_rows=[("T1", "T2", "GHOST", "a", "LEFT JOIN", None, None)],
    )
    rc = main(["generate", "--epic", "E"])
    assert rc == 2
    rejected = edir / "contracts" / "rejected" / "joins.yaml"
    assert rejected.exists()
    payload = yaml.safe_load(rejected.read_text(encoding="utf-8"))
    assert any(e["kind"] == "unknown_join_column" for e in payload["errors"])
    # Canonical joins.yaml is absent.
    assert not (edir / "contracts" / "joins.yaml").exists()


def test_cli_joins_rejection_when_table_missing(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_joins_epic(
        tmp_path, repo_root,
        joins_rows=[("T1", "GHOST", "a", "a", "LEFT JOIN", None, None)],
    )
    rc = main(["generate", "--epic", "E"])
    assert rc == 2
    rejected = edir / "contracts" / "rejected" / "joins.yaml"
    assert rejected.exists()
    payload = yaml.safe_load(rejected.read_text(encoding="utf-8"))
    assert any(e["kind"] == "unknown_join_table" for e in payload["errors"])


# ---------------------------------------------------------------------------
# 9. JOIN_TYPE_ALIASES sanity
# ---------------------------------------------------------------------------


def test_cardinality_custom_separator_accepts_only_declared():
    """When the cardinality column declares `separator: "->"`, only that
    divider parses successfully; other styles return None."""
    assert parse_cardinality("1->n", separator="->") == "1:n"
    assert parse_cardinality("1 -> n", separator="->") == "1:n"  # whitespace OK
    assert parse_cardinality("1:n", separator="->") is None
    assert parse_cardinality("1 to n", separator="->") is None


def test_cardinality_custom_separator_with_spaces():
    """Multi-char separators with spaces work too."""
    assert parse_cardinality("1  -->  n", separator="-->") == "1:n"
    assert parse_cardinality("1->n", separator="-->") is None


def test_cardinality_separator_flows_from_config_through_reader():
    """End-to-end: defaults declares `separator: "->"`, the parser rejects
    `1:n` style values from the spec but accepts `1 -> n`."""
    spec_strict = _spec("""
sheet_name: Joins
column_mapping:
  source_table:  { spec_name: Source Table }
  target_table:  { spec_name: Target Table }
  source_column: { spec_name: Source Col }
  target_column: { spec_name: Target Col }
  join_type:     { spec_name: Type }
  cardinality:   { spec_name: Card,         default_value: null, separator: "->" }
""")
    # Spec uses the declared separator -> parses cleanly.
    wb_ok = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", "LEFT", "1 -> n")])
    res = read_joins_sheet(wb_ok, spec_strict)
    assert not res.errors
    assert res.rows[0].cardinality == "1:n"

    # Spec uses a different separator -> invalid_cardinality.
    wb_bad = _wb_with_joins([("PROJECT", "PROJWBS", "x", "x", "LEFT", "1:n")])
    res2 = read_joins_sheet(wb_bad, spec_strict)
    assert any(e.kind == "invalid_cardinality" for e in res2.errors)


def test_join_type_aliases_has_canonical_set():
    canonical = {"INNER", "LEFT", "RIGHT", "FULL", "CROSS"}
    assert set(JOIN_TYPE_ALIASES.values()) == canonical
