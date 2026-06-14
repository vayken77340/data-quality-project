from pathlib import Path

import yaml

from data_contract.cli import main

from .conftest import add_keys_sheet, minimal_defaults_yaml


def _add_keys_sheet(wb, rows: list[tuple[str, str, str | None]]) -> None:
    """Append a `Keys` sheet using the minimal-defaults schema (ASCII headers)."""
    add_keys_sheet(wb, [(t, pk, fk) for (t, pk, fk) in rows])


def test_generate_epic_1118(tmp_path: Path, repo_root: Path, monkeypatch):
    """Golden-path integration test against a sanitized COPY of the sample spec.

    We don't run against the in-repo spec directly because the source xlsx is
    expected to evolve over time. Instead we mirror the epic into tmp_path,
    patch any blank Obligatoire cells to 'NON' so the run is deterministic,
    and exercise the full generate flow.
    """
    from openpyxl import load_workbook
    from shutil import copy2

    # Mirror configs/types.yaml.
    (tmp_path / "configs").mkdir(parents=True)
    copy2(repo_root / "configs" / "types.yaml", tmp_path / "configs" / "types.yaml")

    # Mirror epic 1118 layout into tmp_path. We use a minimal in-test defaults
    # (sheet `Keys` with ASCII headers) rather than the real defaults.yaml so
    # this test stays decoupled from upstream sheet-naming changes in the real
    # spec workbook.
    edir = tmp_path / "epics" / "1118"
    (edir / "configs").mkdir(parents=True)
    (edir / "specs").mkdir(parents=True)
    (edir / "contracts").mkdir(parents=True)
    (tmp_path / "configs" / "specs_parsing.yaml").write_text(minimal_defaults_yaml(), encoding="utf-8")
    (edir / "configs" / "v1.0.yaml").write_text(
        "epic: 1118\nversion: '1.0'\nspec_file_name: Spec_example.xlsx\n"
        "tables:\n  - table_name: PROJECT\n",
        encoding="utf-8",
    )
    spec_dst = edir / "specs" / "Spec_example.xlsx"
    copy2(repo_root / "epics" / "1118" / "specs" / "Spec_example.xlsx", spec_dst)

    # Patch any blank Obligatoire cell on a non-empty data row to 'NON', so the
    # test is robust to upstream gaps in the sample spec.
    wb = load_workbook(spec_dst)
    ws = wb["PROJECT"]
    obligatoire_col = None
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value and "obligatoire" in str(cell.value).lower():
            obligatoire_col = cell.column
            break
    assert obligatoire_col is not None, "spec sheet must have an Obligatoire column"
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=2).value is None:
            continue  # truly empty row
        if ws.cell(row=row_idx, column=obligatoire_col).value in (None, ""):
            ws.cell(row=row_idx, column=obligatoire_col).value = "NON"

    # Synthesize a `Keys` sheet matching the minimal defaults. Declares proj_id
    # as the PK for PROJECT, no FK.
    _add_keys_sheet(wb, [("PROJECT", "proj_id", None)])
    wb.save(spec_dst)

    monkeypatch.chdir(tmp_path)
    rc = main(["generate", "--epic", "1118"])
    assert rc == 0, "expected a clean run"

    canonical = edir / "contracts" / "PROJECT.yaml"
    history = edir / "contracts" / "history" / "1.0" / "PROJECT.yaml"
    assert canonical.exists()
    assert history.exists()
    assert canonical.read_text(encoding="utf-8") == history.read_text(encoding="utf-8")

    # The source spec is intentionally NOT snapshotted into history/ — every
    # version's spec already lives in specs/ on disk.
    assert not (edir / "contracts" / "history" / "1.0" / "spec").exists()

    # Data dictionary should be auto-emitted as XLSX next to the contracts.
    dd_path = edir / "docs" / "data_dictionary.xlsx"
    assert dd_path.exists()
    from openpyxl import load_workbook
    dd_wb = load_workbook(dd_path, read_only=True)
    assert "README" in dd_wb.sheetnames
    assert "PROJECT" in dd_wb.sheetnames
    assert "Mapping" not in dd_wb.sheetnames
    dd_wb.close()

    data = yaml.safe_load(canonical.read_text(encoding="utf-8"))
    assert data["version"] == "1.0"
    assert data["epic"] == "1118"
    assert data["table"] == "PROJECT"
    assert data["source"]["spec_sheet"] == "PROJECT"
    assert len(data["fields"]) == 5

    by_name = {f["name"]: f for f in data["fields"]}
    assert by_name["proj_id"]["type"] == "float64"
    assert by_name["proj_id"]["nullable"] is False  # OUI -> value_required -> not nullable
    assert by_name["proj_id"]["primary_key"] is True  # from the synthesized Keys sheet
    assert by_name["column1"]["type"] == "boolean"
    # column1's Obligatoire cell in the in-repo spec evolved over time; we
    # only assert presence of the key, not the boolean value.
    assert "nullable" in by_name["column1"]
    assert "primary_key" not in by_name["column1"]
    assert by_name["column2"]["type"] == "timestamp"
    assert by_name["column3"]["type"] == "string"
    assert by_name["column3"]["max_length"] == 384
    assert by_name["column4"]["max_length"] == 384


def test_unknown_type_rejects_table(tmp_path: Path, repo_root: Path, monkeypatch):
    """Inject a temp epic whose spec uses an unknown type; assert it produces a rejected/ file
    and that contracts/<table>.yaml is absent (or deleted)."""
    monkeypatch.chdir(tmp_path)

    epic_root = tmp_path / "epics"
    epic_dir = epic_root / "BAD"
    (epic_dir / "configs").mkdir(parents=True)
    (epic_dir / "specs").mkdir(parents=True)
    (epic_dir / "contracts").mkdir(parents=True)

    # Reuse the global type registry.
    (tmp_path / "configs").mkdir(parents=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo_root / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    # Defaults
    (tmp_path / "configs" / "specs_parsing.yaml").write_text(minimal_defaults_yaml(), encoding="utf-8")

    # Version config
    (epic_dir / "configs" / "v1.0.yaml").write_text(
        "epic: BAD\nversion: '1.0'\nspec_file_name: spec.xlsx\ntables:\n  - table_name: WIDGETS\n",
        encoding="utf-8",
    )

    # Build a tiny xlsx with one bad type.
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "WIDGETS"
    ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    ws.append(["widget_id", "Quaternion(8)", "Identifier", "OUI"])
    _add_keys_sheet(wb, [("WIDGETS", "widget_id", None)])
    wb.save(epic_dir / "specs" / "spec.xlsx")

    rc = main(["generate", "--epic", "BAD", "--epic-root", "epics"])
    assert rc == 2  # rejection → CI signal

    rejected_file = epic_dir / "contracts" / "rejected" / "WIDGETS.yaml"
    canonical_file = epic_dir / "contracts" / "WIDGETS.yaml"
    assert rejected_file.exists()
    assert not canonical_file.exists()
    rej = yaml.safe_load(rejected_file.read_text(encoding="utf-8"))
    assert rej["table"] == "WIDGETS"
    assert any(e["kind"] == "unknown_type" for e in rej["errors"])


def test_generate_all_epics_when_no_epic_arg(tmp_path: Path, repo_root: Path, monkeypatch):
    """No --epic provided: process every epic dir under epics/."""
    monkeypatch.chdir(tmp_path)

    # Mirror the global type registry.
    (tmp_path / "configs").mkdir(parents=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo_root / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    # Build two epics with one tiny spec each.
    from openpyxl import Workbook
    defaults_text = minimal_defaults_yaml()

    for epic in ("100", "200"):
        edir = tmp_path / "epics" / epic
        (edir / "configs").mkdir(parents=True)
        (edir / "specs").mkdir(parents=True)
        (edir / "contracts").mkdir(parents=True)
        (tmp_path / "configs" / "specs_parsing.yaml").write_text(defaults_text, encoding="utf-8")
        (edir / "configs" / "v1.0.yaml").write_text(
            f"epic: '{epic}'\nversion: '1.0'\nspec_file_name: spec.xlsx\n"
            f"tables:\n  - table_name: T\n",
            encoding="utf-8",
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "T"
        ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
        ws.append([f"col_{epic}", "Double", f"desc{epic}", "OUI"])
        _add_keys_sheet(wb, [("T", f"col_{epic}", None)])
        wb.save(edir / "specs" / "spec.xlsx")

    rc = main(["generate"])  # no --epic
    assert rc == 0

    for epic in ("100", "200"):
        assert (tmp_path / "epics" / epic / "contracts" / "T.yaml").exists()
        assert (tmp_path / "epics" / epic / "contracts" / "history" / "1.0" / "T.yaml").exists()


def test_generate_all_epics_aggregates_exit_code(tmp_path: Path, repo_root: Path, monkeypatch):
    """One epic clean + one epic with unknown_type rejection → exit 2."""
    monkeypatch.chdir(tmp_path)

    (tmp_path / "configs").mkdir(parents=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo_root / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    from openpyxl import Workbook
    defaults_text = minimal_defaults_yaml()

    def _make_epic(name: str, type_value: str) -> None:
        edir = tmp_path / "epics" / name
        (edir / "configs").mkdir(parents=True)
        (edir / "specs").mkdir(parents=True)
        (edir / "contracts").mkdir(parents=True)
        (tmp_path / "configs" / "specs_parsing.yaml").write_text(defaults_text, encoding="utf-8")
        (edir / "configs" / "v1.0.yaml").write_text(
            f"epic: '{name}'\nversion: '1.0'\nspec_file_name: spec.xlsx\n"
            f"tables:\n  - table_name: T\n",
            encoding="utf-8",
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "T"
        ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
        ws.append(["a", type_value, "d", "OUI"])
        _add_keys_sheet(wb, [("T", "a", None)])
        wb.save(edir / "specs" / "spec.xlsx")

    _make_epic("GOOD", "Double")
    _make_epic("BAD", "Quaternion(8)")

    rc = main(["generate"])
    assert rc == 2  # at least one rejection
    assert (tmp_path / "epics" / "GOOD" / "contracts" / "T.yaml").exists()
    assert (tmp_path / "epics" / "BAD" / "contracts" / "rejected" / "T.yaml").exists()


def test_config_without_epic_is_an_error(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "configs").mkdir(parents=True)
    (tmp_path / "configs" / "types.yaml").write_text("mappings: []\n", encoding="utf-8")
    rc = main(["generate", "--config", "some.yaml"])
    assert rc == 1


def _bootstrap_multi_version_epic(tmp_path: Path, repo_root: Path, *, versions: list[str], spec_per_version: bool = False) -> Path:
    """Build a self-contained epic at tmp_path/epics/E with the requested version configs.

    If spec_per_version is False (default), all configs share a single Spec.xlsx.
    Otherwise each version gets Spec_v<X>.xlsx.
    """
    (tmp_path / "configs").mkdir(parents=True, exist_ok=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo_root / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    edir = tmp_path / "epics" / "E"
    (edir / "configs").mkdir(parents=True, exist_ok=True)
    (edir / "specs").mkdir(parents=True, exist_ok=True)
    (edir / "contracts").mkdir(parents=True, exist_ok=True)
    (tmp_path / "configs" / "specs_parsing.yaml").write_text(minimal_defaults_yaml(), encoding="utf-8")

    from openpyxl import Workbook

    def _write_spec(name: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "T"
        ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
        ws.append(["a", "Double", "d", "OUI"])
        _add_keys_sheet(wb, [("T", "a", None)])
        wb.save(edir / "specs" / name)

    if spec_per_version:
        for v in versions:
            _write_spec(f"Spec_v{v}.xlsx")
    else:
        _write_spec("Spec.xlsx")

    for v in versions:
        spec_name = f"Spec_v{v}.xlsx" if spec_per_version else "Spec.xlsx"
        (edir / "configs" / f"v{v}.yaml").write_text(
            f"epic: E\nversion: '{v}'\nspec_file_name: {spec_name}\n"
            f"tables:\n  - table_name: T\n",
            encoding="utf-8",
        )
    return edir


def test_backfill_creates_missing_older_history(tmp_path: Path, repo_root: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_multi_version_epic(tmp_path, repo_root, versions=["1.0", "2.0", "3.0"])

    rc = main(["generate", "--epic", "E", "--version", "3.0"])
    assert rc == 0

    # The target version writes canonical + history/3.0
    assert (edir / "contracts" / "T.yaml").exists()
    for v in ("1.0", "2.0", "3.0"):
        assert (edir / "contracts" / "history" / v / "T.yaml").exists(), f"missing v{v}"


def test_backfill_skips_existing_history(tmp_path: Path, repo_root: Path, monkeypatch):
    """If history/v1.0.yaml already exists, the backfill must not overwrite it."""
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_multi_version_epic(tmp_path, repo_root, versions=["1.0", "2.0"])

    sentinel = edir / "contracts" / "history" / "1.0" / "T.yaml"
    sentinel.parent.mkdir(parents=True)
    sentinel.write_text("preserved: true\n", encoding="utf-8")

    rc = main(["generate", "--epic", "E", "--version", "2.0"])
    assert rc == 0
    assert sentinel.read_text(encoding="utf-8") == "preserved: true\n"
    assert (edir / "contracts" / "history" / "2.0" / "T.yaml").exists()


def test_backfill_disabled_via_flag(tmp_path: Path, repo_root: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_multi_version_epic(tmp_path, repo_root, versions=["1.0", "2.0"])

    rc = main(["generate", "--epic", "E", "--version", "2.0", "--no-backfill"])
    assert rc == 0
    assert (edir / "contracts" / "history" / "2.0" / "T.yaml").exists()
    assert not (edir / "contracts" / "history" / "1.0" / "T.yaml").exists()


def test_backfill_skips_when_older_spec_missing(tmp_path: Path, repo_root: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_multi_version_epic(
        tmp_path, repo_root, versions=["1.0", "2.0"], spec_per_version=True
    )
    # Delete the v1.0 spec to simulate the case where the older spec is gone.
    (edir / "specs" / "Spec_v1.0.xlsx").unlink()

    rc = main(["generate", "--epic", "E", "--version", "2.0"])
    assert rc == 0  # target still succeeds
    assert (edir / "contracts" / "history" / "2.0" / "T.yaml").exists()
    assert not (edir / "contracts" / "history" / "1.0" / "T.yaml").exists()


def test_backfill_does_not_overwrite_canonical_or_rejected(tmp_path: Path, repo_root: Path, monkeypatch):
    """Backfill writes only to history/. Canonical and rejected files stay tied to the latest run."""
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_multi_version_epic(tmp_path, repo_root, versions=["1.0", "2.0"])

    # Pre-seed a stale rejected file; backfilling v1.0 must NOT delete it,
    # because backfill never touches canonical/rejected — only the *target* run does.
    rejected_dir = edir / "contracts" / "rejected"
    rejected_dir.mkdir(parents=True)
    stale_rejected = rejected_dir / "OTHER_TABLE.yaml"
    stale_rejected.write_text("untouched: true\n", encoding="utf-8")

    rc = main(["generate", "--epic", "E", "--version", "2.0"])
    assert rc == 0
    assert stale_rejected.read_text(encoding="utf-8") == "untouched: true\n"
