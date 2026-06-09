"""Coverage for the centralized keys sheet (PK + FK enrichment).

Three layers:

1. Sheet-lookup / header / row-parsing layer  -> exercises read_keys_sheet.
2. PK index + enrichment layer                -> exercises build_pk_index +
                                                  enrich_field_contract_list.
3. Configuration parsing                      -> exercises KeysSpec / KeysColumnMapping.

Plus a few CLI-level integration tests for the end-to-end keys flow.
"""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from data_contract.cli import main
from data_contract.config import Defaults, KeysSpec
from data_contract.contract import FieldContract
from data_contract.errors import ConfigError
from data_contract.keys import (
    KeysRow,
    build_pk_index,
    enrich_field_contract_list,
    read_keys_sheet,
    split_separated,
)
from data_contract.spec_reader import open_workbook
from data_contract.type_mapping import Type

from .conftest import add_keys_sheet, minimal_defaults_yaml, minimal_keys_block_yaml


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


_DEFAULT_KEYS_SPEC_YAML = """
sheet_name: Keys
column_mapping:
  table_name:  { spec_name: Table, value_required: true }
  primary_key: { spec_name: PK, value_required: true, separator: "|" }
  foreign_key: { spec_name: FK, value_required: false, separator: "|" }
  comments:    { spec_name: Comments, value_required: false }
"""


def _spec_default() -> KeysSpec:
    return KeysSpec.from_dict(yaml.safe_load(_DEFAULT_KEYS_SPEC_YAML))


def _spec(yaml_str: str) -> KeysSpec:
    return KeysSpec.from_dict(yaml.safe_load(yaml_str))


def _wb_with_keys(rows: list[tuple[str, str, str | None, str | None]], *, sheet_name: str = "Keys") -> Workbook:
    wb = Workbook()
    wb.active.title = "Other"  # leave a non-keys sheet around
    add_keys_sheet(wb, rows, sheet_name=sheet_name)
    return wb


def _field(name: str, type_: Type = Type.STRING, *, nullable: bool = False) -> FieldContract:
    """Default `nullable=False` so PK-enriched fields don't trip the
    `nullable_primary_key` rule unless a test explicitly sets it."""
    return FieldContract(name=name, type=type_, nullable=nullable, description=None)


# ---------------------------------------------------------------------------
# 1. Sheet-lookup layer
# ---------------------------------------------------------------------------


def test_keys_sheet_not_found():
    wb = Workbook()
    wb.active.title = "Only"  # no Keys sheet at all
    result = read_keys_sheet(wb, _spec_default())
    assert len(result.errors) == 1
    assert result.errors[0].kind == "keys_sheet_not_found"


def test_keys_sheet_ambiguous():
    wb = Workbook()
    wb.active.title = "Keys"
    other = wb.create_sheet("keys ")  # trailing space — normalizes to "keys"
    other.append(["Table", "PK"])
    # Both should normalize to "keys"
    result = read_keys_sheet(wb, _spec_default())
    assert any(e.kind == "keys_sheet_ambiguous" for e in result.errors)


def test_keys_sheet_case_insensitive_match():
    wb = Workbook()
    wb.active.title = "KEYS"
    ws = wb.active
    ws.append(["Table", "PK", "FK", "Comments"])
    ws.append(["T", "x", None, None])
    # Spec asks for "Keys"; the workbook has "KEYS" — should match.
    result = read_keys_sheet(wb, _spec_default())
    assert not result.errors
    assert len(result.rows) == 1
    assert result.rows[0].table_name == "T"


# ---------------------------------------------------------------------------
# 1b. Header layer
# ---------------------------------------------------------------------------


def test_keys_header_missing_required_column():
    wb = Workbook()
    ws = wb.create_sheet("Keys")
    ws.append(["Table", "Comments"])  # missing PK column
    ws.append(["T", "anything"])
    result = read_keys_sheet(wb, _spec_default())
    assert any(e.kind == "header_not_found" for e in result.errors)


def test_keys_header_on_third_row():
    wb = Workbook()
    ws = wb.create_sheet("Keys")
    ws.append(["preamble", None, None, None])
    ws.append([None, None, None, None])
    ws.append(["Table", "PK", "FK", "Comments"])
    ws.append(["T", "x", None, None])
    result = read_keys_sheet(wb, _spec_default())
    assert not result.errors
    assert result.rows[0].sheet_row == 4


def test_keys_optional_columns_absent():
    wb = Workbook()
    ws = wb.create_sheet("Keys")
    ws.append(["Table", "PK"])  # no FK column, no Comments
    ws.append(["T", "x"])
    spec = _spec("""
sheet_name: Keys
column_mapping:
  table_name:  { spec_name: Table, value_required: true }
  primary_key: { spec_name: PK, value_required: true, separator: "|" }
""")
    result = read_keys_sheet(wb, spec)
    assert not result.errors
    assert result.rows[0].foreign_keys == []


# ---------------------------------------------------------------------------
# 1c. Row parsing
# ---------------------------------------------------------------------------


def test_keys_row_missing_mandatory_table_name():
    wb = _wb_with_keys([(None, "x", None, None)])  # type: ignore[list-item]
    result = read_keys_sheet(wb, _spec_default())
    assert any(e.kind == "missing_mandatory" and e.field == "table_name" for e in result.errors)


def test_keys_row_missing_mandatory_primary_key():
    wb = _wb_with_keys([("T", None, None, None)])  # type: ignore[list-item]
    result = read_keys_sheet(wb, _spec_default())
    assert any(e.kind == "missing_mandatory" and e.field == "primary_key" for e in result.errors)


def test_keys_row_blank_optional_fk_ok():
    wb = _wb_with_keys([("T", "x", None, None)])
    result = read_keys_sheet(wb, _spec_default())
    assert not result.errors
    assert result.rows[0].foreign_keys == []


def test_keys_row_blank_comments_ok():
    wb = _wb_with_keys([("T", "x", None, None)])
    result = read_keys_sheet(wb, _spec_default())
    assert not result.errors


# ---------------------------------------------------------------------------
# 1d. Separator + whitespace
# ---------------------------------------------------------------------------


def test_pk_separator_split_with_whitespace():
    assert split_separated("  user_id |  region_id  ", "|") == ["user_id", "region_id"]


def test_fk_separator_split_with_whitespace():
    assert split_separated(" col_a|col_b ", "|") == ["col_a", "col_b"]


def test_separator_split_drops_empty_pieces():
    assert split_separated("user_id||region_id|", "|") == ["user_id", "region_id"]


def test_custom_separator():
    spec = _spec("""
sheet_name: Keys
column_mapping:
  table_name:  { spec_name: Table, value_required: true }
  primary_key: { spec_name: PK, value_required: true, separator: "," }
""")
    wb = Workbook()
    ws = wb.create_sheet("Keys")
    ws.append(["Table", "PK"])
    ws.append(["ORDERS", "a, b , c"])
    result = read_keys_sheet(wb, spec)
    assert not result.errors
    assert result.rows[0].primary_keys == ["a", "b", "c"]


# ---------------------------------------------------------------------------
# 2. PK index
# ---------------------------------------------------------------------------


def test_build_pk_index_single_table():
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    idx = build_pk_index(rows)
    assert idx == {"user_id": {"USERS"}}


def test_build_pk_index_composite_pk():
    rows = [KeysRow(2, "ORDER_ITEMS", ["order_id", "item_id"], [])]
    idx = build_pk_index(rows)
    assert idx == {"order_id": {"ORDER_ITEMS"}, "item_id": {"ORDER_ITEMS"}}


def test_build_pk_index_same_column_two_tables():
    rows = [
        KeysRow(2, "TABLE_A", ["id"], []),
        KeysRow(3, "TABLE_B", ["id"], []),
    ]
    idx = build_pk_index(rows)
    assert idx == {"id": {"TABLE_A", "TABLE_B"}}


# ---------------------------------------------------------------------------
# 3. Enrichment
# ---------------------------------------------------------------------------


def test_enrich_pk_only_table():
    fields = [_field("user_id"), _field("name")]
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "USERS", rows, pk_index)
    assert not errors
    assert fields[0].primary_key is True
    assert fields[1].primary_key is None


def test_enrich_composite_pk():
    fields = [_field("order_id"), _field("item_id"), _field("qty")]
    rows = [KeysRow(2, "ORDER_ITEMS", ["order_id", "item_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "ORDER_ITEMS", rows, pk_index)
    assert not errors
    assert fields[0].primary_key is True
    assert fields[1].primary_key is True
    assert fields[2].primary_key is None


def test_enrich_unknown_pk_field():
    fields = [_field("name")]  # no `user_id` field on this table
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "USERS", rows, pk_index)
    assert any(e.kind == "unknown_pk_field" for e in errors)


def test_enrich_fk_resolves_via_pk_index():
    fields = [_field("order_id"), _field("user_id")]
    rows = [
        KeysRow(2, "USERS", ["user_id"], []),
        KeysRow(3, "ORDERS", ["order_id"], ["user_id"]),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "ORDERS", [r for r in rows if r.table_name == "ORDERS"], pk_index
    )
    assert not errors
    assert fields[1].foreign_key == {"table": "USERS", "column": "user_id"}


def test_enrich_unknown_foreign_key_target():
    fields = [_field("ghost_col")]
    rows = [KeysRow(2, "ORDERS", ["order_id"], ["ghost_col"])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "ORDERS", rows, pk_index)
    assert any(e.kind == "unknown_foreign_key_target" for e in errors)


def test_enrich_ambiguous_foreign_key_target():
    fields = [_field("id")]
    rows = [
        KeysRow(2, "TABLE_A", ["id"], []),
        KeysRow(3, "TABLE_B", ["id"], []),
        KeysRow(4, "ORDERS",  ["order_id"], ["id"]),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "ORDERS", [r for r in rows if r.table_name == "ORDERS"], pk_index
    )
    assert any(e.kind == "ambiguous_foreign_key_target" for e in errors)


def test_enrich_fk_column_missing_from_table():
    fields = [_field("user_id")]  # FK refers to legacy_id which doesn't exist on this table
    rows = [
        KeysRow(2, "USERS",  ["user_id"], []),
        KeysRow(3, "ORDERS", ["order_id"], ["legacy_id"]),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "ORDERS", [r for r in rows if r.table_name == "ORDERS"], pk_index
    )
    assert any(e.kind == "unknown_foreign_key_target" for e in errors)


def test_enrich_multi_fk_per_row():
    fields = [_field("order_id"), _field("user_id"), _field("product_id"), _field("qty")]
    rows = [
        KeysRow(2, "USERS",    ["user_id"],    []),
        KeysRow(3, "PRODUCTS", ["product_id"], []),
        KeysRow(4, "ORDERS",   ["order_id"],   ["user_id", "product_id"]),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "ORDERS", [r for r in rows if r.table_name == "ORDERS"], pk_index
    )
    assert not errors
    assert fields[1].foreign_key == {"table": "USERS", "column": "user_id"}
    assert fields[2].foreign_key == {"table": "PRODUCTS", "column": "product_id"}
    assert fields[3].foreign_key is None


def test_enrich_skips_tables_not_being_generated():
    """Rows for LEGACY_X are in the keys sheet, but we only enrich USERS+ORDERS.
    No error from the LEGACY_X row."""
    fields_users = [_field("user_id")]
    rows = [
        KeysRow(2, "USERS",    ["user_id"], []),
        KeysRow(3, "LEGACY_X", ["junk"], []),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields_users, "USERS", [r for r in rows if r.table_name == "USERS"], pk_index
    )
    assert not errors
    assert fields_users[0].primary_key is True


def test_enrich_multiple_keys_rows_same_table_merge():
    fields = [_field("user_id"), _field("region_id"), _field("group_id")]
    rows = [
        KeysRow(2, "USERS", ["user_id"], []),
        KeysRow(3, "USERS", ["region_id"], ["group_id"]),
        KeysRow(4, "GROUPS", ["group_id"], []),
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "USERS", [r for r in rows if r.table_name == "USERS"], pk_index
    )
    assert not errors
    assert fields[0].primary_key is True
    assert fields[1].primary_key is True
    assert fields[2].foreign_key == {"table": "GROUPS", "column": "group_id"}


def test_enrich_pk_with_nullable_true_rejects():
    """A field flagged as PK by the keys sheet must not be nullable.
    Otherwise: nullable_primary_key rejection."""
    fields = [FieldContract(name="user_id", type=Type.INT64, nullable=True, description=None)]
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "USERS", rows, pk_index)
    assert any(e.kind == "nullable_primary_key" and e.field == "user_id" for e in errors)


def test_enrich_pk_with_nullable_false_ok():
    fields = [FieldContract(name="user_id", type=Type.INT64, nullable=False, description=None)]
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "USERS", rows, pk_index)
    assert errors == []
    assert fields[0].primary_key is True


def test_enrich_pk_with_nullable_none_ok():
    """nullable=None means `value_required: false` on the source column with a blank
    cell — no explicit declaration. The PK+nullable rule only fires on
    nullable=True (explicit `is nullable`), not on absence."""
    fields = [FieldContract(name="user_id", type=Type.INT64, nullable=None, description=None)]
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(fields, "USERS", rows, pk_index)
    assert errors == []


def test_enrich_fk_allow_violations_demotes_unknown_target():
    """FK rule violations land in `fk_warnings`, not `errors`, when
    allow_violations=True. The field's foreign_key annotation just isn't set."""
    fields = [_field("order_id"), _field("ghost_col")]
    rows = [KeysRow(2, "ORDERS", ["order_id"], ["ghost_col"])]
    pk_index = build_pk_index(rows)
    _, errors, warnings = enrich_field_contract_list(
        fields, "ORDERS", rows, pk_index, fk_allow_violations=True,
    )
    assert errors == []
    assert any(w.kind == "unknown_foreign_key_target" for w in warnings)
    assert fields[1].foreign_key is None


def test_enrich_fk_allow_violations_demotes_ambiguous_target():
    fields = [_field("order_id"), _field("id")]
    rows = [
        KeysRow(2, "TABLE_A", ["id"], []),
        KeysRow(3, "TABLE_B", ["id"], []),
        KeysRow(4, "ORDERS",  ["order_id"], ["id"]),
    ]
    pk_index = build_pk_index(rows)
    _, errors, warnings = enrich_field_contract_list(
        fields, "ORDERS",
        [r for r in rows if r.table_name == "ORDERS"],
        pk_index,
        fk_allow_violations=True,
    )
    assert errors == []
    assert any(w.kind == "ambiguous_foreign_key_target" for w in warnings)


def test_enrich_fk_allow_violations_does_not_demote_pk_errors():
    """PK errors are NOT downgraded by allow_violations (it's FK-specific)."""
    fields = [_field("name")]  # no `user_id` field — PK reference is broken
    rows = [KeysRow(2, "USERS", ["user_id"], [])]
    pk_index = build_pk_index(rows)
    _, errors, warnings = enrich_field_contract_list(
        fields, "USERS", rows, pk_index, fk_allow_violations=True,
    )
    assert any(e.kind == "unknown_pk_field" for e in errors)
    assert warnings == []


def test_enrich_duplicate_pk_declaration_is_idempotent():
    fields = [_field("user_id")]
    rows = [
        KeysRow(2, "USERS", ["user_id"], []),
        KeysRow(3, "USERS", ["user_id"], []),  # duplicate
    ]
    pk_index = build_pk_index(rows)
    _, errors, _ = enrich_field_contract_list(
        fields, "USERS", [r for r in rows if r.table_name == "USERS"], pk_index
    )
    assert not errors
    assert fields[0].primary_key is True


# ---------------------------------------------------------------------------
# 4. Configuration parsing
# ---------------------------------------------------------------------------


_VALID_COLUMN_MAPPING_YAML = """
fields:
  column_mapping:
    name:        { spec_name: N, value_required: true }
    type:        { spec_name: T, value_required: true }
    description: { spec_name: D, value_required: false }
    nullable:
      spec_name: Obligatoire
      value_required: true
      values:
        "true":  ["non"]
        "false": ["oui"]
"""


def test_keys_block_missing_in_defaults_raises_config_error(tmp_path):
    p = tmp_path / "defaults.yaml"
    p.write_text(_VALID_COLUMN_MAPPING_YAML, encoding="utf-8")
    with pytest.raises(ConfigError, match="keys"):
        Defaults.from_yaml(p)


def test_keys_block_missing_sheet_name_raises_config_error(tmp_path):
    p = tmp_path / "defaults.yaml"
    p.write_text(_VALID_COLUMN_MAPPING_YAML + """
keys:
  column_mapping:
    table_name:  { spec_name: Table, value_required: true }
    primary_key: { spec_name: PK, value_required: true, separator: "|" }
""", encoding="utf-8")
    with pytest.raises(ConfigError, match="sheet_name"):
        Defaults.from_yaml(p)


def test_keys_block_partial_columns_only_required_two():
    spec = _spec("""
sheet_name: Keys
column_mapping:
  table_name:  { spec_name: Table, value_required: true }
  primary_key: { spec_name: PK, value_required: true, separator: "|" }
""")
    assert spec.column_mapping.foreign_key is None
    assert spec.column_mapping.comments is None


def test_keys_separator_required_non_empty():
    with pytest.raises(ConfigError, match="separator"):
        _spec("""
sheet_name: Keys
column_mapping:
  table_name:  { spec_name: Table, value_required: true }
  primary_key: { spec_name: PK, value_required: true, separator: "" }
""")


# ---------------------------------------------------------------------------
# 5. FieldContract round-trip
# ---------------------------------------------------------------------------


def test_field_contract_round_trip_with_foreign_key():
    f = FieldContract(
        name="user_id",
        type=Type.INT64,
        nullable=False,
        description="ref",
        primary_key=True,
        foreign_key={"table": "USERS", "column": "user_id"},
    )
    payload = f.to_dict()
    f2 = FieldContract.from_dict(payload)
    assert f2.primary_key is True
    assert f2.foreign_key == {"table": "USERS", "column": "user_id"}
    assert f2.constraints == {}  # PK + FK are not on the constraints dict


def test_contract_from_dict_picks_up_keys_block_fields():
    payload = {
        "name": "user_id",
        "type": "integer",
        "primary_key": True,
        "foreign_key": {"table": "USERS", "column": "user_id"},
        "some_custom_constraint": "value",
    }
    f = FieldContract.from_dict(payload)
    assert f.primary_key is True
    assert f.foreign_key == {"table": "USERS", "column": "user_id"}
    assert f.constraints == {"some_custom_constraint": "value"}


# ---------------------------------------------------------------------------
# 6. CLI integration: end-to-end keys flow + new rejection kinds.
# ---------------------------------------------------------------------------


def _bootstrap_keys_epic(
    tmp_path: Path,
    repo_root: Path,
    *,
    tables_in_keys: list[tuple[str, str, str | None]] | None = None,
    skip_keys_sheet: bool = False,
    epic_tables: list[str] | None = None,
) -> Path:
    """Build a fresh tmp epic with one 'T' table-structure sheet and a Keys
    sheet according to `tables_in_keys`. If `skip_keys_sheet=True`, leave the
    Keys sheet out entirely."""
    (tmp_path / "configs").mkdir(parents=True, exist_ok=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    edir = tmp_path / "epics" / "E"
    (edir / "configs").mkdir(parents=True)
    (edir / "specs").mkdir(parents=True)
    (edir / "contracts").mkdir(parents=True)

    epic_tables = epic_tables or ["T"]
    tables_in_keys = tables_in_keys if tables_in_keys is not None else [("T", "x", None)]

    # Defaults: declares the standard column_mapping plus a Keys block.
    (edir / "configs" / "defaults.yaml").write_text(
        """
fields:
  column_mapping:
    name:        { spec_name: Champ dans extract, value_required: true }
    type:        { spec_name: Type, value_required: true }
    description: { spec_name: Description, value_required: false }
    nullable:
      spec_name: Obligatoire
      value_required: true
      values:
        "true":  ["non"]
        "false": ["oui"]
""" + minimal_keys_block_yaml(),
        encoding="utf-8",
    )
    table_lines = "\n".join(f"  - table_name: {t}" for t in epic_tables)
    (edir / "configs" / "v1.0.yaml").write_text(
        f"epic: E\nversion: '1.0'\nspec_file_name: spec.xlsx\ntables:\n{table_lines}\n",
        encoding="utf-8",
    )

    wb = Workbook()
    wb.active.title = epic_tables[0]
    for t in epic_tables:
        ws = wb[t] if t in wb.sheetnames else wb.create_sheet(t)
        ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
        ws.append(["x", "Double", "d", "OUI"])
    if not skip_keys_sheet:
        add_keys_sheet(wb, tables_in_keys)
    wb.save(edir / "specs" / "spec.xlsx")
    return edir


def test_cli_happy_path_pk_emits_primary_key_flag(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_keys_epic(tmp_path, repo_root)
    rc = main(["generate", "--epic", "E"])
    assert rc == 0
    data = yaml.safe_load((edir / "contracts" / "T.yaml").read_text(encoding="utf-8"))
    by_name = {f["name"]: f for f in data["fields"]}
    assert by_name["x"]["primary_key"] is True


def test_cli_happy_path_fk_emits_structured_foreign_key(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_keys_epic(
        tmp_path, repo_root,
        epic_tables=["USERS", "ORDERS"],
        tables_in_keys=[
            ("USERS",  "user_id", None),
            ("ORDERS", "order_id", "user_id"),
        ],
    )
    # Adjust the workbook so each table has its referenced columns.
    from openpyxl import load_workbook
    wb = load_workbook(edir / "specs" / "spec.xlsx")
    wb["USERS"].delete_rows(1, wb["USERS"].max_row)
    wb["USERS"].append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    wb["USERS"].append(["user_id", "Double", "d", "OUI"])
    wb["ORDERS"].delete_rows(1, wb["ORDERS"].max_row)
    wb["ORDERS"].append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    wb["ORDERS"].append(["order_id", "Double", "d", "OUI"])
    wb["ORDERS"].append(["user_id",  "Double", "d", "NON"])
    wb.save(edir / "specs" / "spec.xlsx")

    rc = main(["generate", "--epic", "E"])
    assert rc == 0
    orders = yaml.safe_load((edir / "contracts" / "ORDERS.yaml").read_text(encoding="utf-8"))
    by_name = {f["name"]: f for f in orders["fields"]}
    assert by_name["user_id"]["foreign_key"] == {"table": "USERS", "column": "user_id"}


def test_cli_keys_sheet_not_found_rejects_every_table(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_keys_epic(tmp_path, repo_root, skip_keys_sheet=True)
    rc = main(["generate", "--epic", "E"])
    assert rc == 2
    rejected = edir / "contracts" / "rejected" / "T.yaml"
    assert rejected.exists()
    rej = yaml.safe_load(rejected.read_text(encoding="utf-8"))
    assert any(e["kind"] == "keys_sheet_not_found" for e in rej["errors"])


def test_cli_keys_missing_table_rejects_that_table_only(tmp_path, repo_root, monkeypatch):
    """Epic generates USERS + ORDERS. The keys sheet has a row only for USERS.
    ORDERS gets rejected with keys_missing_table; USERS is built normally."""
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_keys_epic(
        tmp_path, repo_root,
        epic_tables=["USERS", "ORDERS"],
        tables_in_keys=[("USERS", "x", None)],
    )
    rc = main(["generate", "--epic", "E"])
    assert rc == 2
    assert (edir / "contracts" / "USERS.yaml").exists()
    rejected = edir / "contracts" / "rejected" / "ORDERS.yaml"
    assert rejected.exists()
    rej = yaml.safe_load(rejected.read_text(encoding="utf-8"))
    assert any(e["kind"] == "keys_missing_table" for e in rej["errors"])


def test_cli_fk_allow_violations_builds_clean(tmp_path, repo_root, monkeypatch):
    """With `allow_foreign_key_violation=true` from the project `.env` (the
    in-code default when `.env` is absent), an FK row pointing at a column
    that doesn't exist anywhere is demoted to a warning. The contract still
    builds (exit 0) and the offending field has no foreign_key."""
    monkeypatch.chdir(tmp_path)
    (tmp_path / "configs").mkdir(parents=True, exist_ok=True)
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo_root / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    edir = tmp_path / "epics" / "E"
    (edir / "configs").mkdir(parents=True)
    (edir / "specs").mkdir(parents=True)
    (edir / "contracts").mkdir(parents=True)

    (edir / "configs" / "defaults.yaml").write_text(minimal_defaults_yaml(), encoding="utf-8")
    (edir / "configs" / "v1.0.yaml").write_text(
        "epic: E\nversion: '1.0'\nspec_file_name: spec.xlsx\n"
        "tables:\n  - table_name: T\n",
        encoding="utf-8",
    )

    wb = Workbook()
    wb.active.title = "T"
    ws = wb["T"]
    ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    ws.append(["x", "Double", "d", "OUI"])
    # Keys sheet declares `ghost` as an FK on table T — ghost doesn't exist as
    # a field. With allow_foreign_key_violation=true, that's a warning not a rejection.
    add_keys_sheet(wb, [("T", "x", "ghost")])
    wb.save(edir / "specs" / "spec.xlsx")

    rc = main(["generate", "--epic", "E"])
    assert rc == 0
    canonical = edir / "contracts" / "T.yaml"
    assert canonical.exists()
    data = yaml.safe_load(canonical.read_text(encoding="utf-8"))
    by_name = {f["name"]: f for f in data["fields"]}
    # `x` got its PK flag; no `foreign_key` because the FK target was unresolved
    # and allow_foreign_key_violation downgraded it to a warning.
    assert by_name["x"].get("primary_key") is True
    assert "foreign_key" not in by_name["x"]


def test_cli_unknown_pk_field_rejection(tmp_path, repo_root, monkeypatch):
    monkeypatch.chdir(tmp_path)
    edir = _bootstrap_keys_epic(
        tmp_path, repo_root,
        tables_in_keys=[("T", "no_such_column", None)],
    )
    rc = main(["generate", "--epic", "E"])
    assert rc == 2
    rej = yaml.safe_load((edir / "contracts" / "rejected" / "T.yaml").read_text(encoding="utf-8"))
    assert any(e["kind"] == "unknown_pk_field" for e in rej["errors"])
