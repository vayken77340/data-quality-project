from pathlib import Path
from shutil import copy2

from openpyxl import Workbook

from data_contract.cli import main

from .conftest import add_keys_sheet, minimal_defaults_yaml


def _bootstrap_epic(tmp_path: Path, repo_root: Path, *, bad_type: bool = False) -> Path:
    (tmp_path / "configs").mkdir(parents=True)
    copy2(repo_root / "configs" / "types.yaml", tmp_path / "configs" / "types.yaml")
    edir = tmp_path / "epics" / "E"
    (edir / "configs").mkdir(parents=True)
    (edir / "specs").mkdir(parents=True)
    (edir / "contracts").mkdir(parents=True)
    (tmp_path / "configs" / "specs_parsing.yaml").write_text(minimal_defaults_yaml(), encoding="utf-8")
    (edir / "configs" / "v1.0.yaml").write_text(
        "epic: E\nversion: '1.0'\nspec_file_name: spec.xlsx\n"
        "tables:\n  - table_name: T\n",
        encoding="utf-8",
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws.append(["Champ dans extract", "Type", "Description", "Obligatoire ?"])
    ws.append(["x", "Quaternion(8)" if bad_type else "Double", "d", "OUI"])
    add_keys_sheet(wb, [("T", "x")])
    wb.save(edir / "specs" / "spec.xlsx")
    return edir


def _listing(p: Path) -> set[Path]:
    if not p.exists():
        return set()
    return {f for f in p.rglob("*") if f.is_file()}


def test_lint_clean_exits_0_and_writes_nothing(tmp_path: Path, repo_root: Path, monkeypatch):
    edir = _bootstrap_epic(tmp_path, repo_root)
    before = _listing(edir / "contracts")
    monkeypatch.chdir(tmp_path)
    rc = main(["lint", "--epic", "E"])
    after = _listing(edir / "contracts")
    assert rc == 0
    assert after == before


def test_lint_unknown_type_exits_2_and_writes_nothing(tmp_path: Path, repo_root: Path, monkeypatch):
    edir = _bootstrap_epic(tmp_path, repo_root, bad_type=True)
    before = _listing(edir / "contracts")
    monkeypatch.chdir(tmp_path)
    rc = main(["lint", "--epic", "E"])
    after = _listing(edir / "contracts")
    assert rc == 2
    assert after == before


def test_generate_after_lint_produces_artifacts(tmp_path: Path, repo_root: Path, monkeypatch, capsys):
    """Sanity: lint is a true no-op and generate still works."""
    edir = _bootstrap_epic(tmp_path, repo_root)
    monkeypatch.chdir(tmp_path)
    assert main(["lint", "--epic", "E"]) == 0
    assert main(["generate", "--epic", "E"]) == 0
    assert (edir / "contracts" / "T.yaml").exists()
    assert (edir / "contracts" / "history" / "1.0" / "T.yaml").exists()


def test_lint_prefix_in_output(tmp_path: Path, repo_root: Path, monkeypatch, capsys):
    _bootstrap_epic(tmp_path, repo_root)
    monkeypatch.chdir(tmp_path)
    main(["lint", "--epic", "E"])
    out = capsys.readouterr().out
    assert "[LINT-OK]" in out


def test_lint_rejection_prefix(tmp_path: Path, repo_root: Path, monkeypatch, capsys):
    _bootstrap_epic(tmp_path, repo_root, bad_type=True)
    monkeypatch.chdir(tmp_path)
    main(["lint", "--epic", "E"])
    captured = capsys.readouterr()
    assert "[LINT-REJECTED]" in (captured.out + captured.err)


def test_drift_command_between_two_history_snapshots(tmp_path: Path, repo_root: Path, monkeypatch):
    edir = _bootstrap_epic(tmp_path, repo_root)
    # Build v1.0
    monkeypatch.chdir(tmp_path)
    assert main(["generate", "--epic", "E"]) == 0
    # Hand-author a v2.0 snapshot by editing the spec and re-generating under v2.0.
    (edir / "configs" / "v1.0.yaml").unlink()
    (edir / "configs" / "v2.0.yaml").write_text(
        "epic: E\nversion: '2.0'\nspec_file_name: spec.xlsx\n"
        "tables:\n  - table_name: T\n",
        encoding="utf-8",
    )
    # Add a column to the spec so v2.0 differs from v1.0.
    from openpyxl import load_workbook
    wb = load_workbook(edir / "specs" / "spec.xlsx")
    ws = wb["T"]
    ws.append(["y", "Double", "d", "NON"])
    wb.save(edir / "specs" / "spec.xlsx")
    assert main(["generate", "--epic", "E", "--version", "2.0"]) == 0

    # Standalone drift command (no rebuild)
    rc = main(["drift", "--epic", "E", "--table", "T", "--from", "1.0", "--to", "2.0", "--write"])
    assert rc == 0
    drift_file = edir / "contracts" / "drift" / "T__v1.0_to_v2.0.yaml"
    assert drift_file.exists()
