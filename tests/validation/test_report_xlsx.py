"""Tests for the business-facing XLSX report."""

from __future__ import annotations

from pathlib import Path
from textwrap import dedent

import yaml
from openpyxl import load_workbook

from data_contract.cli import main
from tests.conftest import ALL_CHECKS_ENABLED_YAML


def _build_epic(tmp_path: Path, *, target=None) -> Path:
    epic = tmp_path / "epics" / "T"
    (epic / "configs").mkdir(parents=True)
    (epic / "contracts").mkdir()
    (epic / "sample").mkdir()
    repo = Path(__file__).resolve().parents[2]
    (tmp_path / "configs").mkdir()
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "targets").mkdir()
    for n in ("oracle.yaml", "postgres.yaml", "iceberg.yaml"):
        (tmp_path / "configs" / "targets" / n).write_text(
            (repo / "configs" / "targets" / n).read_text(encoding="utf-8"),
            encoding="utf-8",
        )

    contract = {
        "version": "1.0", "epic": "T", "table": "T",
        "source": {"spec_file": "s", "spec_sheet": "T"},
        "fields": [
            {"name": "id", "type": "int64", "nullable": False, "primary_key": True},
            {"name": "label", "type": "string", "nullable": False, "max_length": 3},
        ],
    }
    (epic / "contracts" / "T.yaml").write_text(
        yaml.safe_dump(contract, sort_keys=False), encoding="utf-8",
    )
    target_line = f"target: {target or 'postgres'}\n"
    (epic / "configs" / "validation.yaml").write_text(
        ALL_CHECKS_ENABLED_YAML + target_line +
        dedent("""\
            defaults:
              format: csv
              file_pattern: "sample/{table}.csv"
        """),
        encoding="utf-8",
    )
    return epic


def _run(tmp_path: Path, csv_content: str, *, target=None) -> Path:
    epic = _build_epic(tmp_path, target=target)
    (epic / "sample" / "T.csv").write_text(csv_content, encoding="utf-8")
    out = tmp_path / "out"
    main([
        "validate-data", "--epic", "T",
        "--epic-root", str(tmp_path / "epics"),
        "--input-dir", str(epic),
        "--output-dir", str(out),
        "--types", str(tmp_path / "configs" / "types.yaml"),
    ])
    return out / "quality_report.xlsx"


def test_xlsx_sheets_for_clean_run(tmp_path: Path):
    p = _run(tmp_path, "id,label\n1,abc\n2,def\n")
    wb = load_workbook(p, read_only=True)
    names = wb.sheetnames
    # Always-present sheets in the gold-standard layout:
    assert "Run" in names
    assert "Summary" in names
    assert "Profile" in names              # consolidated per-table profile sheet
    assert "Checks" in names
    # The standalone Dimensions sheet was merged into Summary.
    assert "Dimensions" not in names
    assert "Hints" not in names
    # No per-table profile sheets; everything is on the global Profile sheet.
    assert "T_profile" not in names
    # No rejected sheet on a clean run.
    assert "T_rejected" not in names
    assert "T_rejected_rows" not in names
    assert "T_rejected_violations" not in names
    # No Table issues sheet on a clean run.
    assert "Run issues" not in names
    assert "Table issues" not in names
    wb.close()


def test_xlsx_run_sheet_target_when_set(tmp_path: Path):
    p = _run(tmp_path, "id,label\n1,abc\n", target="postgres")
    wb = load_workbook(p, read_only=True)
    run_ws = wb["Run"]
    # Find the Target row.
    target_row = None
    for row in run_ws.iter_rows(values_only=True):
        if row and row[0] == "Target":
            target_row = row
            break
    assert target_row is not None
    assert target_row[1] == "postgres"
    wb.close()


def test_xlsx_summary_sheet_score(tmp_path: Path):
    p = _run(tmp_path, "id,label\n1,abc\n")
    wb = load_workbook(p, read_only=True)
    summary = wb["Summary"]
    rows = list(summary.iter_rows(values_only=True))
    # First row is "Overall clean row %" + value.
    assert rows[0][0] == "Overall clean row %"
    assert rows[0][1] == "100.0"
    # Status row.
    assert rows[1][0] == "Status"
    assert rows[1][1] == "PASS"
    wb.close()


def test_xlsx_summary_inlines_dimension_columns(tmp_path: Path):
    """The standalone Dimensions sheet is gone; per-table dimension scores
    sit inline as columns on the right of the Summary scorecard."""
    p = _run(tmp_path, "id,label\n1,abc\n")
    wb = load_workbook(p, read_only=True)
    summary = wb["Summary"]
    rows = list(summary.iter_rows(values_only=True))
    scorecard_header = next(r for r in rows if r and r[0] == "Table")
    for dim in ("Completeness", "PK Uniqueness", "FK Consistency"):
        assert dim in scorecard_header, f"missing dimension column {dim!r}"
    # Dropped / renamed columns:
    for dropped in ("Validity", "Uniqueness", "Consistency", "Score", "Pass Rate"):
        assert dropped not in scorecard_header, f"unexpected column {dropped!r}"
    assert "Info" in scorecard_header
    assert "Clean Row %" in scorecard_header
    # Total Rows moved to AFTER Clean Rows.
    clean_idx = scorecard_header.index("Clean Rows")
    total_idx = scorecard_header.index("Total Rows")
    assert total_idx == clean_idx + 1, (
        f"Total Rows should sit right after Clean Rows; "
        f"got Clean Rows@{clean_idx}, Total Rows@{total_idx}"
    )
    wb.close()


def test_xlsx_summary_totals_row_present(tmp_path: Path):
    """A `Total` row at the bottom of the scorecard sums Errors / Warnings /
    Info / Files / Rows and shows the overall clean-row %."""
    p = _run(tmp_path, "id,label\n1,toolong\n2,abc\n")
    wb = load_workbook(p, read_only=True)
    summary = wb["Summary"]
    rows = list(summary.iter_rows(values_only=True))
    last_row = next(r for r in reversed(rows) if r and r[0] is not None)
    assert last_row[0] == "Total"
    # Columns: Table | PK | Files | Clean Rows | Total Rows | Errors | Warnings | Info | Clean Row % | dims...
    assert last_row[2] == 1                  # one file
    assert last_row[3] == 1                  # one clean row
    assert last_row[4] == 2                  # two total rows
    assert last_row[5] == 1                  # one error
    wb.close()


def test_xlsx_profile_sheet_has_table_section(tmp_path: Path):
    """The unified Profile sheet stacks one section per table, each preceded
    by a `Table: <name>` heading row."""
    p = _run(tmp_path, "id,label\n1,abc\n")
    wb = load_workbook(p, read_only=True)
    profile = wb["Profile"]
    rows = list(profile.iter_rows(values_only=True))
    # First non-empty row is the heading for the T section.
    assert rows[0][0] == "Table: T"
    # Second row is the profile column header row.
    assert rows[1][0] == "Field"
    wb.close()


def test_xlsx_rejected_sheet_when_violations(tmp_path: Path):
    """One row per violation. Header layout:
        Severity | Source File | <PK fields> | <only fields with a violation> | Check | Expected.
    Severity values are uppercased; only the Severity cell is colour-filled."""
    p = _run(tmp_path, "id,label\n1,toolong\n")
    wb = load_workbook(p)
    assert "T_rejected" in wb.sheetnames
    assert "T_rejected_rows" not in wb.sheetnames
    assert "T_rejected_violations" not in wb.sheetnames

    rj = wb["T_rejected"]
    headers = [rj.cell(row=1, column=c).value for c in range(1, rj.max_column + 1)]
    # Severity is now the first column.
    assert headers[0] == "Severity"
    assert headers[1] == "Source File"
    assert headers[2] == "id"            # PK sits right after Source File
    assert "Worst Severity" not in headers
    assert "Source Row" not in headers
    # Only fields with a violation appear in the contract-field block.
    # The CSV has a max_length violation on `label`; `id` has no violation,
    # so it should NOT appear as a field column (the PK column at index 2 is
    # the only `id` header).
    assert headers.count("id") == 1      # PK only, no field-block duplicate
    assert "label" in headers
    # Right-hand annotation columns.
    assert headers[-2] == "Check"
    assert headers[-1] == "Expected"

    # Severity value uppercased.
    assert rj.cell(row=2, column=1).value == "ERROR"
    # Only the Severity cell is colour-filled; Source File and other cells
    # stay uncoloured.
    severity_fill = rj.cell(row=2, column=1).fill.fgColor.rgb
    source_file_fill = rj.cell(row=2, column=2).fill.fgColor.rgb
    assert severity_fill not in (None, "00000000")
    assert source_file_fill in (None, "00000000")

    # The label cell carries the offending value; Check is business-friendly.
    label_col = headers.index("label") + 1
    assert rj.cell(row=2, column=label_col).value == "toolong"
    assert rj.cell(row=2, column=len(headers) - 1).value == "Value too long"
    assert rj.cell(row=2, column=len(headers)).value
    wb.close()


def test_xlsx_rejected_sheet_one_row_per_source_row_stacked_cells(tmp_path: Path):
    """A source row with multiple violations becomes ONE sheet row; the Check
    and Expected cells stack one line per violation so a reader can map check
    N to expected N at the same vertical offset."""
    # max_length violation on `label` AND nullable violation on `code`.
    epic = tmp_path / "epics" / "M"
    (epic / "configs").mkdir(parents=True)
    (epic / "contracts").mkdir()
    (epic / "sample").mkdir()
    repo = Path(__file__).resolve().parents[2]
    (tmp_path / "configs").mkdir()
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "targets").mkdir()
    for n in ("oracle.yaml", "postgres.yaml", "iceberg.yaml"):
        (tmp_path / "configs" / "targets" / n).write_text(
            (repo / "configs" / "targets" / n).read_text(encoding="utf-8"),
            encoding="utf-8",
        )
    contract = {
        "version": "1.0", "epic": "M", "table": "M",
        "source": {"spec_file": "s", "spec_sheet": "M"},
        "fields": [
            {"name": "id", "type": "int64", "nullable": False, "primary_key": True},
            {"name": "label", "type": "string", "nullable": False, "max_length": 3},
            {"name": "code", "type": "string", "nullable": False},
        ],
    }
    (epic / "contracts" / "M.yaml").write_text(
        yaml.safe_dump(contract, sort_keys=False), encoding="utf-8",
    )
    (epic / "configs" / "validation.yaml").write_text(
        ALL_CHECKS_ENABLED_YAML + "target: postgres\n" + dedent("""\
            defaults:
              format: csv
              file_pattern: "sample/{table}.csv"
        """),
        encoding="utf-8",
    )
    (epic / "sample" / "M.csv").write_text("id,label,code\n1,toolong,\n", encoding="utf-8")
    out = tmp_path / "out"
    main([
        "validate-data", "--epic", "M",
        "--epic-root", str(tmp_path / "epics"),
        "--input-dir", str(epic),
        "--output-dir", str(out),
        "--types", str(tmp_path / "configs" / "types.yaml"),
    ])
    wb = load_workbook(out / "quality_report.xlsx", read_only=True)
    rj = wb["M_rejected"]
    headers = [rj.cell(row=1, column=c).value for c in range(1, rj.max_column + 1)]
    # Layout: Severity | Source File | id (PK) | <violating fields> | Check | Expected.
    assert headers[0] == "Severity"
    assert headers[1] == "Source File"
    assert headers[2] == "id"            # PK
    assert headers[-2] == "Check"
    assert headers[-1] == "Expected"

    # Two violations on the same source row -> ONE sheet row (stacked cells).
    data_rows = [tuple(rj.cell(row=r, column=c).value for c in range(1, rj.max_column + 1))
                 for r in range(2, rj.max_row + 1)]
    assert len(data_rows) == 1

    # Only 'label' and 'code' have violations; 'id' has no violation so it
    # does NOT appear as a field column (only the PK `id` is present).
    assert "label" in headers
    assert "code" in headers
    assert headers.count("id") == 1      # PK only

    label_col = headers.index("label")
    code_col = headers.index("code")
    check_col = len(headers) - 2
    expected_col = len(headers) - 1
    only = data_rows[0]
    # Field-block cells carry the offending value for each violating field.
    assert only[code_col] == "(null)"
    assert only[label_col] == "toolong"
    # Check and Expected are newline-stacked, sorted: errors first, then by
    # field name -> 'code' before 'label'.
    assert only[check_col] == "Missing value\nValue too long"
    assert only[expected_col] == "required\nmax 3 chars"
    # Severity = worst across the row, uppercased.
    assert only[0] == "ERROR"
    wb.close()


def test_xlsx_rejected_sheet_uses_source_row_when_no_pk(tmp_path: Path):
    """Without PK fields, the slot right after Source File is `Source Row`."""
    # Build an epic whose contract has no primary_key on any field.
    epic = tmp_path / "epics" / "NoPK"
    (epic / "configs").mkdir(parents=True)
    (epic / "contracts").mkdir()
    (epic / "sample").mkdir()
    repo = Path(__file__).resolve().parents[2]
    (tmp_path / "configs").mkdir()
    (tmp_path / "configs" / "types.yaml").write_text(
        (repo / "configs" / "types.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "parsers.yaml").write_text(
        (repo / "configs" / "parsers.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "configs" / "targets").mkdir()
    for n in ("oracle.yaml", "postgres.yaml", "iceberg.yaml"):
        (tmp_path / "configs" / "targets" / n).write_text(
            (repo / "configs" / "targets" / n).read_text(encoding="utf-8"),
            encoding="utf-8",
        )
    contract = {
        "version": "1.0", "epic": "NoPK", "table": "NoPK",
        "source": {"spec_file": "s", "spec_sheet": "NoPK"},
        "fields": [
            {"name": "label", "type": "string", "nullable": False, "max_length": 3},
        ],
    }
    (epic / "contracts" / "NoPK.yaml").write_text(
        yaml.safe_dump(contract, sort_keys=False), encoding="utf-8",
    )
    (epic / "configs" / "validation.yaml").write_text(
        ALL_CHECKS_ENABLED_YAML + "target: postgres\n" + dedent("""\
            defaults:
              format: csv
              file_pattern: "sample/{table}.csv"
        """),
        encoding="utf-8",
    )
    (epic / "sample" / "NoPK.csv").write_text("label\ntoolong\n", encoding="utf-8")
    out = tmp_path / "out"
    main([
        "validate-data", "--epic", "NoPK",
        "--epic-root", str(tmp_path / "epics"),
        "--input-dir", str(epic),
        "--output-dir", str(out),
        "--types", str(tmp_path / "configs" / "types.yaml"),
    ])
    wb = load_workbook(out / "quality_report.xlsx", read_only=True)
    assert "NoPK_rejected" in wb.sheetnames
    rj = wb["NoPK_rejected"]
    headers = [rj.cell(row=1, column=c).value for c in range(1, rj.max_column + 1)]
    assert headers[0] == "Severity"
    assert headers[1] == "Source File"
    assert headers[2] == "Source Row"  # no PK -> Source Row sits in the slot
    wb.close()


def test_xlsx_table_issues_sheet_surfaces_column_missing(tmp_path: Path):
    """`column_missing` is a table-level violation (no source row). It must
    surface in the Table issues sheet with severity tint, friendly check label,
    field name, terse expected, and hint -- otherwise it's invisible in Excel.
    """
    # CSV missing the contract's `label` column.
    p = _run(tmp_path, "id\n1\n")
    wb = load_workbook(p, read_only=True)
    assert "Table issues" in wb.sheetnames
    sh = wb["Table issues"]
    rows = list(sh.iter_rows(values_only=True))
    assert rows[0] == ("Severity", "Table", "Check", "Field", "Expected", "Hint")
    # Find the column_missing row for `label`.
    body = [r for r in rows[1:] if r[3] == "label"]
    assert len(body) == 1
    severity, table, check, field, expected, hint = body[0]
    assert severity == "ERROR"
    assert table == "T"
    assert check == "Column missing in source"
    assert field == "label"
    assert "label" in expected               # carries the missing column name
    assert hint                              # non-empty action sentence
    wb.close()


def test_xlsx_checks_sheet_lists_enabled_only(tmp_path: Path):
    """The Checks sheet (renamed from Hints) lists only ENABLED checks with
    their YAML-supplied description and the emitted violation kind."""
    p = _run(tmp_path, "id,label\n1,abc\n")
    wb = load_workbook(p, read_only=True)
    checks = wb["Checks"]
    rows = list(checks.iter_rows(values_only=True))
    assert rows[0] == ("Check", "Description", "Violation kind")
    # Every row past the header is an enabled check; description and kind populated.
    check_names = [r[0] for r in rows[1:]]
    assert "type_coercion" in check_names
    assert "pk_uniqueness" in check_names
    # Every row carries a non-empty description and violation kind.
    for r in rows[1:]:
        assert r[1] and r[1].strip(), f"empty description for {r[0]}"
        assert r[2] and r[2].strip(), f"empty violation kind for {r[0]}"
    wb.close()
